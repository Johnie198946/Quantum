from importlib.metadata import version
assert version('openai')=='2.24.0'
import os,tempfile,socket,threading,time,io,json,statistics,math
from pathlib import Path
root=Path(tempfile.mkdtemp(prefix='travel-office-perf-'))
with socket.socket() as sock:
 sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
os.environ['AI_LAB_HOME']=str(root/'vault');os.environ['HERMES_TENANT_SANDBOX_ROOT']=str(root/'sandboxes')
os.environ['HERMES_BRIDGE_INTERNAL_TOKEN']='isolated-office-performance';os.environ['HERMES_BRIDGE_URL']=f'http://127.0.0.1:{port}/v1/chat'
from scripts import hermes_bridge as b
from backend.services.document_sources import save_document_source,document_text,document_original_path,DocumentSourceError
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Inches
import uvicorn
server=uvicorn.Server(uvicorn.Config(b.app,host='127.0.0.1',port=port,lifespan='off',log_level='error'))
threading.Thread(target=server.run,daemon=True).start()
for _ in range(500):
 if server.started:break
 time.sleep(.02)
assert server.started
pdf=Image.new('RGB',(1000,600),'white');draw=ImageDraw.Draw(pdf);draw.text((100,100),'Booking: SCAN-2048\nCheck-in: 2030-01-03',fill='black',font_size=45)
pdf_bytes=io.BytesIO();pdf.save(pdf_bytes,'PDF')
pictures=[]
for i in range(16):
 im=Image.new('RGB',(1000,500),'white');draw=ImageDraw.Draw(im);draw.rectangle((50,50,180,250),fill=(i*15,80,200));draw.text((230,100),f'Booking: TEST-{i:03d}',fill='black',font_size=45)
 out=io.BytesIO();im.save(out,'PNG');pictures.append(out.getvalue())
ppt=Presentation()
for i in range(40):
 slide=ppt.slides.add_slide(ppt.slide_layouts[6]);slide.shapes.add_picture(io.BytesIO(pictures[i%16]),Inches(.5),Inches(.5),width=Inches(9))
 slide.shapes.add_textbox(Inches(.5),Inches(5.2),Inches(8),Inches(1)).text_frame.text='Confirmed travel reference page '+str(i+1)
ppt_bytes=io.BytesIO();ppt.save(ppt_bytes)
rows=[]
try:
 for kind,data,expected in [('scanned_pdf',pdf_bytes.getvalue(),['SCAN-2048','2030']),('40_slide_16_image_pptx',ppt_bytes.getvalue(),[f'TEST-{i:03d}' for i in range(16)])]:
  suffix='.pdf' if kind=='scanned_pdf' else '.pptx'
  for i in range(30):
   start=time.monotonic()
   r=save_document_source(tenant_key='office-performance',user_id='owner',filename='reference'+suffix,content_type='application/octet-stream',data=data)
   elapsed=time.monotonic()-start
   assert r['status']=='ready',r.get('parse_error')
   text,_=document_text('office-performance','owner',r['source_id']);assert all(x in text for x in expected),(kind,'visual_fields_missing')
   original,_=document_original_path('office-performance','owner',r['source_id']);assert original.read_bytes()==data
   try: document_original_path('office-performance','other',r['source_id']);raise AssertionError('cross_owner_access')
   except DocumentSourceError as e: assert e.code=='document_not_found'
   row={'kind':kind,'sample':i,'cache':'cold' if i==0 else 'warm','seconds':round(elapsed,3),'passed':True,'original_retained':True,'note_created':bool(r.get('note_id')),'source_bytes':len(data)}
   rows.append(row);print(json.dumps(row),flush=True)
  values=sorted(r['seconds'] for r in rows if r['kind']==kind)
  print(json.dumps({'summary':kind,'n':len(values),'cold_n':1,'warm_n':29,'p50_s':round(statistics.median(values),3),'p95_s':values[math.ceil(.95*len(values))-1],'max_s':max(values),'scope':'isolated local document service -> real Hermes vision -> private note; excludes upload network, contribution compile and client UI; one cold sample is not cold p95'}),flush=True)
finally:server.should_exit=True
