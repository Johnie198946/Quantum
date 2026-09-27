from importlib.metadata import version
assert version("openai") == "2.24.0", "Use requirements-bridge-worker.lock runtime (OpenAI 2.24.0)"
import os,tempfile,socket,threading,time,io,json
from pathlib import Path
root=Path(tempfile.mkdtemp(prefix='travel-office-upload-'))
with socket.socket() as sock:
 sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
os.environ['AI_LAB_HOME']=str(root/'vault');os.environ['HERMES_TENANT_SANDBOX_ROOT']=str(root/'sandboxes')
os.environ['HERMES_BRIDGE_INTERNAL_TOKEN']='office-local-test-secret';os.environ['HERMES_BRIDGE_URL']=f'http://127.0.0.1:{port}/v1/chat'
from scripts import hermes_bridge as b
from backend.services.document_sources import save_document_source,document_text,document_original_path
from PIL import Image,ImageDraw
from docx import Document
from pptx import Presentation
from pptx.util import Inches
import uvicorn
server=uvicorn.Server(uvicorn.Config(b.app,host='127.0.0.1',port=port,lifespan='off',log_level='error'))
threading.Thread(target=server.run,daemon=True).start()
while not server.started:time.sleep(.02)
im=Image.new('RGB',(900,400),'white');draw=ImageDraw.Draw(im);draw.rectangle((40,40,220,220),fill='blue');draw.text((260,80),'Hotel booking: TEST-2048\nCheck-in: 2030-01-03',fill='black',font_size=35)
image=io.BytesIO();im.save(image,'PNG')
doc=Document();doc.add_paragraph('Office visual reference');doc.add_picture(io.BytesIO(image.getvalue()))
ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[6]);slide.shapes.add_picture(io.BytesIO(image.getvalue()),Inches(1),Inches(1))
try:
 for document,suffix in [(doc,'.docx'),(ppt,'.pptx'),(doc,'.doc'),(ppt,'.ppt')]:
  data=io.BytesIO();document.save(data)
  if suffix in {'.doc','.ppt'}:
   import subprocess
   modern=root/('source'+('.docx' if suffix=='.doc' else '.pptx'));modern.write_bytes(data.getvalue())
   result=subprocess.run(['/Applications/LibreOffice.app/Contents/MacOS/soffice','-env:UserInstallation='+ (root/('lo-profile'+suffix)).as_uri(),'--headless','--convert-to',suffix[1:],'--outdir',str(root),str(modern)],capture_output=True,text=True,timeout=90)
   converted=root/('source'+suffix);assert result.returncode==0 and converted.exists(),result.stderr
   data=io.BytesIO(converted.read_bytes())
  started=time.monotonic()
  receipt=save_document_source(tenant_key='office-upload-test',user_id='alice',filename='booking'+suffix,content_type='application/octet-stream',data=data.getvalue())
  assert receipt['status']=='ready',receipt.get('parse_error')
  text,_=document_text('office-upload-test','alice',receipt['source_id'])
  assert '2048' in text and '2030' in text
  original,_=document_original_path('office-upload-test','alice',receipt['source_id']);assert original.read_bytes()==data.getvalue()
  print(json.dumps({'status':'passed','type':suffix,'seconds':round(time.monotonic()-started,2),'private_note':bool(receipt.get('note_id')),'original_retained':True,'visual_fields_extracted':True,'root':str(root)}),flush=True)
finally:server.should_exit=True
