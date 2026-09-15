"""Deterministic structural and rendered visual gates for PPTX artifacts."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from PIL import Image, ImageChops, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader


class PresentationVisualGateError(ValueError):
    pass


def _layout(slide) -> str:
    markers = [
        shape.name.split(":", 1)[1]
        for shape in slide.shapes
        if shape.name.startswith("layout:")
    ]
    if len(markers) != 1:
        raise PresentationVisualGateError("each slide needs exactly one layout marker")
    return markers[0]


def _text(slide) -> str:
    return " ".join(
        shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    ).strip()


def inspect_pptx(pptx_path: Path) -> dict[str, Any]:
    prs = Presentation(pptx_path)
    slide_width, slide_height = prs.slide_width, prs.slide_height
    layouts, errors, picture_hashes = [], [], []
    vector_icon_ids: set[str] = set()
    map_shape_names: set[str] = set()
    total_shapes = 0
    for slide_index, slide in enumerate(prs.slides, 1):
        try:
            layouts.append(_layout(slide))
        except PresentationVisualGateError as exc:
            errors.append(f"slide {slide_index}: {exc}")
            layouts.append("unknown")
        if not _text(slide):
            errors.append(f"slide {slide_index}: blank text layer")
        meaningful = [
            shape for shape in slide.shapes if not shape.name.startswith("layout:")
        ]
        if not meaningful:
            errors.append(f"slide {slide_index}: blank page")
        text_shapes = []
        for shape in meaningful:
            total_shapes += 1
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > slide_width + 1000
                or shape.top + shape.height > slide_height + 1000
            ):
                errors.append(f"slide {slide_index}: out-of-bounds object {shape.name}")
            if shape.name.startswith("vector:"):
                parts = shape.name.split(":")
                if len(parts) >= 3:
                    vector_icon_ids.add(parts[1])
            if shape.name.startswith("map:"):
                map_shape_names.add(shape.name)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_hashes.append(hashlib.sha256(shape.image.blob).hexdigest())
                crops = [
                    shape.crop_left,
                    shape.crop_right,
                    shape.crop_top,
                    shape.crop_bottom,
                ]
                if any(abs(value) > 1e-8 for value in crops):
                    errors.append(f"slide {slide_index}: cropped picture {shape.name}")
                source_ratio = shape.image.size[0] / shape.image.size[1]
                placed_ratio = shape.width / shape.height
                if abs(source_ratio - placed_ratio) / source_ratio > 0.01:
                    errors.append(
                        f"slide {slide_index}: distorted picture {shape.name}"
                    )
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                exempt = (
                    shape.name.startswith(
                        ("source-caption:", "page-number:", "map:attribution")
                    )
                    or shape.text.strip().isdigit()
                )
                if not exempt:
                    text_shapes.append(shape)
                    sizes = [
                        run.font.size.pt
                        for paragraph in shape.text_frame.paragraphs
                        for run in paragraph.runs
                        if run.text and run.font.size
                    ]
                    if sizes and min(sizes) < 18:
                        errors.append(
                            f"slide {slide_index}: text below 18pt in {shape.name}"
                        )
        for first_index, first in enumerate(text_shapes):
            first_box = (
                first.left,
                first.top,
                first.left + first.width,
                first.top + first.height,
            )
            for second in text_shapes[first_index + 1 :]:
                second_box = (
                    second.left,
                    second.top,
                    second.left + second.width,
                    second.top + second.height,
                )
                overlap_width = min(first_box[2], second_box[2]) - max(
                    first_box[0], second_box[0]
                )
                overlap_height = min(first_box[3], second_box[3]) - max(
                    first_box[1], second_box[1]
                )
                if overlap_width > 1000 and overlap_height > 1000:
                    overlap = overlap_width * overlap_height
                    smaller = min(
                        first.width * first.height, second.width * second.height
                    )
                    if smaller and overlap / smaller > 0.02:
                        errors.append(
                            f"slide {slide_index}: overlapping text objects {first.name} and {second.name}"
                        )
    if len(set(layouts)) < 5:
        errors.append("deck uses fewer than five layout families")
    if any(left == right for left, right in zip(layouts, layouts[1:])):
        errors.append("consecutive slides repeat the same layout family")
    editable_backdrop = {"map:europe-land", "map:asia-land", "map:bosphorus-water"}
    licensed_real_base = {"map:base-real", "map:attribution"}
    if not (
        editable_backdrop <= map_shape_names or licensed_real_base <= map_shape_names
    ):
        errors.append(
            "geographic map is missing an editable backdrop or attributed real base"
        )
    landmark_nodes = {
        name
        for name in map_shape_names
        if name.startswith("map:landmark-") and "label" not in name
    }
    if len(landmark_nodes) < 4:
        errors.append("geographic map has fewer than four editable landmark nodes")
    if len(vector_icon_ids) < 6:
        errors.append("deck contains fewer than six editable SVG-derived icons")
    if len(set(picture_hashes)) < 5:
        errors.append("deck contains fewer than five distinct raster scene assets")
    report = {
        "slide_count": len(prs.slides),
        "layout_families": sorted(set(layouts)),
        "layout_sequence": layouts,
        "shape_count": total_shapes,
        "picture_count": len(picture_hashes),
        "distinct_picture_count": len(set(picture_hashes)),
        "editable_vector_icon_count": len(vector_icon_ids),
        "editable_map_landmark_count": len(landmark_nodes),
        "errors": errors,
    }
    return report


def inspect_pdf(pdf_path: Path, *, expected_pages: int) -> dict[str, Any]:
    reader = PdfReader(str(pdf_path))
    errors = []
    if len(reader.pages) != expected_pages:
        errors.append(
            f"PDF page count {len(reader.pages)} differs from PPTX {expected_pages}"
        )
    for index, page in enumerate(reader.pages, 1):
        box = page.mediabox
        if float(box.width) <= 0 or float(box.height) <= 0:
            errors.append(f"PDF page {index} has invalid dimensions")
        if not (page.extract_text() or "").strip():
            errors.append(f"PDF page {index} has no extractable text")
    return {"page_count": len(reader.pages), "errors": errors}


def render_pdf_pages(
    pdf_path: Path, output_dir: Path, *, scale: float = 1.2
) -> list[Path]:
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise RuntimeError("pypdfium2 is required for rendered-page gates") from exc
    output_dir.mkdir(parents=True, exist_ok=True)
    document = pdfium.PdfDocument(str(pdf_path))
    paths: list[Path] = []
    try:
        for index in range(len(document)):
            page = document[index]
            bitmap = page.render(scale=scale)
            image = bitmap.to_pil().convert("RGB").copy()
            bitmap.close()
            page.close()
            extrema = ImageChops.difference(
                image, Image.new("RGB", image.size, image.getpixel((0, 0)))
            ).getbbox()
            if extrema is None:
                raise PresentationVisualGateError(
                    f"rendered PDF page {index + 1} is blank"
                )
            path = output_dir / f"slide-{index + 1:02d}.png"
            image.save(path, format="PNG", optimize=True)
            paths.append(path)
    finally:
        document.close()
    return paths


def build_montage(
    page_paths: list[Path],
    output_path: Path,
    *,
    columns: int = 3,
    thumb_width: int = 640,
) -> Path:
    if not page_paths:
        raise ValueError("montage needs rendered pages")
    thumbs = []
    for path in page_paths:
        image = Image.open(path).convert("RGB")
        height = round(image.height * thumb_width / image.width)
        thumbs.append(image.resize((thumb_width, height), Image.Resampling.LANCZOS))
    gap, label_height = 28, 34
    rows = (len(thumbs) + columns - 1) // columns
    cell_height = max(image.height for image in thumbs) + label_height
    canvas = Image.new(
        "RGB",
        (
            columns * thumb_width + (columns + 1) * gap,
            rows * cell_height + (rows + 1) * gap,
        ),
        "#E8E5DF",
    )
    draw = ImageDraw.Draw(canvas)
    for index, image in enumerate(thumbs):
        col, row = index % columns, index // columns
        x, y = gap + col * (thumb_width + gap), gap + row * (cell_height + gap)
        canvas.paste(image, (x, y + label_height))
        draw.text((x, y + 4), f"{index + 1:02d}", fill="#172A3A")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output_path, format="PNG", optimize=True)
    return output_path


def run_visual_gates(
    pptx_path: Path,
    pdf_path: Path,
    *,
    report_path: Path | None = None,
    render_dir: Path | None = None,
    montage_path: Path | None = None,
) -> dict[str, Any]:
    pptx = inspect_pptx(pptx_path)
    pdf = inspect_pdf(pdf_path, expected_pages=pptx["slide_count"])
    pages = render_pdf_pages(pdf_path, render_dir) if render_dir else []
    if montage_path:
        build_montage(pages, montage_path)
    report = {
        "pptx": pptx,
        "pdf": pdf,
        "rendered_page_count": len(pages),
        "passed": not (pptx["errors"] or pdf["errors"]),
    }
    if report_path:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    if not report["passed"]:
        raise PresentationVisualGateError("; ".join(pptx["errors"] + pdf["errors"]))
    return report
