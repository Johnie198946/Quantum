"""Bounded editable PPTX generation and honest PDF rendering."""

from __future__ import annotations

import base64
import json
import math
import shutil
import subprocess
import tempfile
import warnings
from io import BytesIO
from pathlib import Path
from typing import Any

from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from backend.services.presentation_map import project_point, validate_geo_points
from backend.services.presentation_materials import validate_material_bytes
from backend.services.presentation_scenario import DEFAULT_THEME, validate_theme
from backend.services.presentation_svg import parse_editable_svg

_LAYOUT_FIELDS = {
    "title": {"layout", "title", "subtitle"},
    "section": {"layout", "title"},
    "bullets": {"layout", "title", "subtitle", "bullets"},
    "conclusion": {"layout", "title", "subtitle", "bullets"},
    "two_column": {"layout", "title", "subtitle", "left", "right"},
    "table": {"layout", "title", "headers", "rows"},
    "chart": {"layout", "title", "categories", "series"},
    "timeline": {"layout", "title", "subtitle", "events"},
    "hero_photo": {"layout", "title", "subtitle", "photo", "caption"},
    "photo_collage": {"layout", "title", "subtitle", "photos"},
    "geo_route_map": {
        "layout",
        "title",
        "subtitle",
        "map",
        "attribution",
        "points",
    },
    "icon_facts": {"layout", "title", "subtitle", "items"},
    "quote_photo": {"layout", "title", "quote", "attribution", "photo"},
    "data_story": {"layout", "title", "subtitle", "metric", "unit", "body", "facts"},
    "icon_grid": {"layout", "title", "subtitle", "items"},
    "route_map": {"layout", "title", "subtitle", "points"},
    "image": {"layout", "title", "subtitle", "image_data", "caption"},
}

_ICON_LABELS = {
    "camera": "CAM",
    "card": "PAY",
    "ferry": "SEA",
    "food": "EAT",
    "hotel": "BED",
    "map": "MAP",
    "shield": "SAFE",
    "train": "RAIL",
    "walk": "WALK",
    "landmark": "SEE",
}


def _validated_material(
    value: Any, *, expected: set[str]
) -> tuple[bytes, dict[str, Any]]:
    if not isinstance(value, dict) or set(value) != {"data_uri", "manifest"}:
        raise ValueError("visual material needs only data_uri and manifest")
    encoded, supplied = value.get("data_uri"), value.get("manifest")
    if (
        not isinstance(encoded, str)
        or not isinstance(supplied, dict)
        or "," not in encoded
    ):
        raise ValueError("visual material data or manifest is invalid")
    header, payload = encoded.split(",", 1)
    mime = {
        "data:image/png;base64": "image/png",
        "data:image/jpeg;base64": "image/jpeg",
        "data:image/svg+xml;base64": "image/svg+xml",
    }.get(header)
    if mime not in expected:
        raise ValueError("visual material MIME is not valid for this layout")
    try:
        raw = base64.b64decode(payload, validate=True)
    except Exception as exc:
        raise ValueError("visual material base64 is invalid") from exc
    verified = validate_material_bytes(
        raw,
        declared_mime=mime,
        source_url=str(supplied.get("source_url") or ""),
        author=str(supplied.get("author") or ""),
        license_id=str(supplied.get("license_id") or ""),
        license_url=str(supplied.get("license_url") or ""),
        fetched_at=str(supplied.get("fetched_at") or ""),
        commercial_use_allowed=supplied.get("commercial_use_allowed") is True,
    )
    comparable = {key: supplied.get(key) for key in verified}
    if comparable != verified or set(supplied) - (set(verified) | {"cache_path"}):
        raise ValueError("visual material manifest does not match decoded bytes")
    if supplied.get("cache_status") not in {"validated", "ready"}:
        raise ValueError("visual material is not validated")
    return raw, verified


def _theme(value: dict[str, Any]) -> dict[str, Any]:
    raw = validate_theme(value.get("theme", DEFAULT_THEME))
    return {
        "colors": {
            key: RGBColor.from_string(color[1:]) for key, color in raw["colors"].items()
        },
        "fonts": raw["fonts"],
    }


def _spec(content: str) -> dict[str, Any]:
    try:
        value = json.loads(content)
    except json.JSONDecodeError as exc:
        raise ValueError("presentation artifact must be valid JSON") from exc
    slides = value.get("slides") if isinstance(value, dict) else None
    if not isinstance(slides, list) or not 1 <= len(slides) <= 60:
        raise ValueError("presentation needs between 1 and 60 slides")
    if set(value) - {"title", "subtitle", "theme", "slides"}:
        raise ValueError("presentation contains unsupported fields")
    total_image_bytes = 0
    for index, slide in enumerate(slides, 1):
        if not isinstance(slide, dict):
            raise ValueError(f"slide {index} is invalid")
        layout = str(slide.get("layout") or "bullets")
        if layout not in _LAYOUT_FIELDS:
            raise ValueError(f"slide {index} has unsupported layout {layout}")
        if set(slide) - _LAYOUT_FIELDS[layout]:
            raise ValueError(f"slide {index} contains fields unused by {layout} layout")
        for key, limit in (("title", 180), ("subtitle", 240)):
            if key in slide and len(str(slide[key])) > limit:
                raise ValueError(f"slide {index} {key} exceeds {limit} characters")
        for key, count, length in (
            ("bullets", 12, 500),
            ("left", 12, 500),
            ("right", 12, 500),
            ("headers", 8, 100),
            ("categories", 20, 80),
        ):
            if key not in slide:
                continue
            items = slide[key]
            if not isinstance(items, list) or len(items) > count:
                raise ValueError(f"slide {index} {key} exceeds {count} items")
            if any(len(str(item)) > length for item in items):
                raise ValueError(
                    f"slide {index} {key} item exceeds {length} characters"
                )
        if layout == "table":
            headers, rows = slide.get("headers"), slide.get("rows")
            if (
                not isinstance(headers, list)
                or not headers
                or not isinstance(rows, list)
                or len(rows) > 12
            ):
                raise ValueError(f"slide {index} table is invalid or exceeds 12 rows")
            if any(
                not isinstance(row, list)
                or len(row) != len(headers)
                or any(len(str(cell)) > 160 for cell in row)
                for row in rows
            ):
                raise ValueError(
                    f"slide {index} table rows must match headers and cells must not exceed 160 characters"
                )
        if layout == "chart":
            categories, series = slide.get("categories"), slide.get("series")
            if (
                not isinstance(categories, list)
                or not categories
                or not isinstance(series, list)
                or not 1 <= len(series) <= 6
            ):
                raise ValueError(f"slide {index} chart is invalid or exceeds 6 series")
            for item in series:
                values = item.get("values") if isinstance(item, dict) else None
                if (
                    not isinstance(values, list)
                    or len(values) != len(categories)
                    or len(str(item.get("name") or "系列")) > 80
                ):
                    raise ValueError(
                        f"slide {index} chart series must match categories"
                    )
                try:
                    if any(not math.isfinite(float(number)) for number in values):
                        raise ValueError
                except (TypeError, ValueError) as exc:
                    raise ValueError(
                        f"slide {index} chart values must be finite numbers"
                    ) from exc
        if layout in {"hero_photo", "quote_photo"}:
            _validated_material(
                slide.get("photo"), expected={"image/png", "image/jpeg"}
            )
        if layout == "photo_collage":
            photos = slide.get("photos")
            if not isinstance(photos, list) or not 2 <= len(photos) <= 4:
                raise ValueError(f"slide {index} photo collage needs 2-4 photos")
            for photo in photos:
                if not isinstance(photo, dict) or set(photo) != {"material", "caption"}:
                    raise ValueError(f"slide {index} photo collage entry is invalid")
                if len(str(photo.get("caption") or "")) > 80:
                    raise ValueError(f"slide {index} photo caption is too long")
                _validated_material(
                    photo.get("material"), expected={"image/png", "image/jpeg"}
                )
        if layout == "icon_facts":
            items = slide.get("items")
            if not isinstance(items, list) or not 2 <= len(items) <= 8:
                raise ValueError(f"slide {index} icon facts needs 2-8 items")
            for item in items:
                if not isinstance(item, dict) or set(item) != {
                    "icon",
                    "title",
                    "detail",
                }:
                    raise ValueError(f"slide {index} icon fact is invalid")
                if (
                    not str(item.get("title") or "").strip()
                    or len(str(item.get("title"))) > 60
                    or len(str(item.get("detail") or "")) > 120
                ):
                    raise ValueError(f"slide {index} icon fact text is invalid")
                raw, _ = _validated_material(
                    item.get("icon"), expected={"image/svg+xml"}
                )
                parse_editable_svg(raw)
        if layout == "geo_route_map":
            validate_geo_points(slide.get("points"))
            if slide.get("map") is not None:
                _validated_material(
                    slide["map"], expected={"image/png", "image/jpeg"}
                )
        if layout == "data_story":
            facts = slide.get("facts")
            if (
                not isinstance(facts, list)
                or not 1 <= len(facts) <= 4
                or any(len(str(fact)) > 120 for fact in facts)
            ):
                raise ValueError(f"slide {index} data story facts are invalid")
            if (
                len(str(slide.get("metric") or "")) > 16
                or len(str(slide.get("body") or "")) > 280
            ):
                raise ValueError(f"slide {index} data story text is too long")
        if layout == "quote_photo":
            if (
                not str(slide.get("quote") or "").strip()
                or len(str(slide.get("quote"))) > 260
            ):
                raise ValueError(f"slide {index} quote is invalid")
        if layout in {"timeline", "icon_grid", "route_map"}:
            key = {"timeline": "events", "icon_grid": "items", "route_map": "points"}[
                layout
            ]
            items = slide.get(key)
            limits = {"timeline": (2, 8), "icon_grid": (2, 8), "route_map": (2, 10)}[
                layout
            ]
            if not isinstance(items, list) or not limits[0] <= len(items) <= limits[1]:
                raise ValueError(
                    f"slide {index} {layout} needs {limits[0]}-{limits[1]} items"
                )
            allowed = {
                "timeline": {"label", "title", "detail"},
                "icon_grid": {"icon", "title", "detail"},
                "route_map": {"name", "side", "detail"},
            }[layout]
            required = {
                "timeline": {"title"},
                "icon_grid": {"icon", "title"},
                "route_map": {"name", "side"},
            }[layout]
            field_limits = {
                "timeline": {"label": 24, "title": 60, "detail": 120},
                "icon_grid": {"icon": 16, "title": 60, "detail": 120},
                "route_map": {"name": 60, "side": 8, "detail": 100},
            }[layout]
            for item in items:
                if not isinstance(item, dict) or set(item) - allowed:
                    raise ValueError(f"slide {index} {layout} item is invalid")
                if any(not str(item.get(field) or "").strip() for field in required):
                    raise ValueError(f"slide {index} {layout} item is incomplete")
                if any(
                    len(str(item.get(field) or "")) > limit
                    for field, limit in field_limits.items()
                ):
                    raise ValueError(f"slide {index} {layout} item is too long")
                if layout == "icon_grid" and item.get("icon") not in _ICON_LABELS:
                    raise ValueError(f"slide {index} icon is unsupported")
                if layout == "route_map" and item.get("side") not in {
                    "europe",
                    "asia",
                    "route",
                }:
                    raise ValueError(f"slide {index} route-map side is unsupported")
            if layout == "route_map" and any(
                sum(item["side"] == side for item in items) > 4
                for side in {"europe", "asia", "route"}
            ):
                raise ValueError(
                    f"slide {index} route-map has too many points per side"
                )
        if layout == "image":
            encoded = str(slide.get("image_data") or "")
            if not encoded.startswith(
                ("data:image/png;base64,", "data:image/jpeg;base64,")
            ):
                raise ValueError(
                    f"slide {index} image must be an embedded PNG or JPEG data URI"
                )
            try:
                raw = base64.b64decode(encoded.split(",", 1)[1], validate=True)
            except Exception as exc:
                raise ValueError(f"slide {index} image data is invalid") from exc
            if not 32 <= len(raw) <= 8_000_000:
                raise ValueError(f"slide {index} image size is invalid")
            total_image_bytes += len(raw)
            if total_image_bytes > 24_000_000:
                raise ValueError("presentation images exceed 24 MB")
            declared_format = "PNG" if encoded.startswith("data:image/png;") else "JPEG"
            expected_magic = (
                b"\x89PNG\r\n\x1a\n" if declared_format == "PNG" else b"\xff\xd8\xff"
            )
            if not raw.startswith(expected_magic):
                raise ValueError(
                    f"slide {index} image signature does not match MIME type"
                )
            try:
                with warnings.catch_warnings():
                    warnings.simplefilter("error", Image.DecompressionBombWarning)
                    with Image.open(BytesIO(raw)) as image:
                        width, height = image.size
                        actual_format = image.format
                        image.verify()
                if actual_format != declared_format:
                    raise ValueError(
                        f"slide {index} image format does not match MIME type"
                    )
                if width < 64 or height < 64 or width * height > 24_000_000:
                    raise ValueError(f"slide {index} image dimensions are invalid")
            except (
                UnidentifiedImageError,
                OSError,
                Image.DecompressionBombWarning,
            ) as exc:
                raise ValueError(
                    f"slide {index} image cannot be decoded safely"
                ) from exc
    return value


def _text(
    shape,
    theme: dict[str, Any],
    size: int = 22,
    color: str = "text",
    bold: bool = False,
    font: str = "body",
) -> None:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = theme["fonts"][font]
            run.font.size = Pt(size)
            run.font.color.rgb = theme["colors"][color]
            run.font.bold = bold


def _title(slide, title: str, theme: dict[str, Any], subtitle: str = "") -> None:
    box = slide.shapes.add_textbox(
        Inches(0.65), Inches(0.42), Inches(11.9), Inches(0.8)
    )
    box.text_frame.text = title
    _text(box, theme, 28, bold=True, font="title")
    line = slide.shapes.add_shape(
        1, Inches(0.65), Inches(1.22), Inches(1.15), Inches(0.08)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = theme["colors"]["primary"]
    line.line.fill.background()
    if subtitle:
        sub = slide.shapes.add_textbox(
            Inches(0.67), Inches(1.4), Inches(11.4), Inches(0.55)
        )
        sub.text_frame.text = subtitle
        _text(sub, theme, 18, "muted")


def _bullets(
    slide,
    items: list[Any],
    theme: dict[str, Any],
    *,
    x: float = 0.8,
    y: float = 1.8,
    w: float = 11.6,
    h: float = 4.8,
) -> None:
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x - 0.18),
        Inches(y - 0.14),
        Inches(w + 0.36),
        Inches(h + 0.22),
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = theme["colors"]["pale"]
    panel.line.fill.background()
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = str(item)
        paragraph.level = 0
        paragraph.space_after = Pt(12)
        paragraph.font.name = theme["fonts"]["body"]
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = theme["colors"]["text"]
    _text(box, theme, 20)


def _table(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    rows = spec.get("rows") or []
    headers = spec.get("headers") or []
    if not isinstance(headers, list) or not headers or not isinstance(rows, list):
        _bullets(slide, spec.get("bullets") or [], theme)
        return
    cols = len(headers)
    body = rows
    table = slide.shapes.add_table(
        len(body) + 1, cols, Inches(0.7), Inches(1.7), Inches(11.9), Inches(4.9)
    ).table
    for col, value in enumerate(headers):
        table.cell(0, col).text = str(value)
    for r, row in enumerate(body, 1):
        for c, value in enumerate(row):
            table.cell(r, c).text = str(value)
    for r in range(len(body) + 1):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = theme["colors"]["primary" if r == 0 else "pale"]
            for p in cell.text_frame.paragraphs:
                p.font.name = theme["fonts"]["body"]
                p.font.size = Pt(13)
                p.font.color.rgb = theme["colors"]["inverse" if r == 0 else "text"]
                for run in p.runs:
                    run.font.name = theme["fonts"]["body"]
                    run.font.size = Pt(13)
                    run.font.color.rgb = theme["colors"][
                        "inverse" if r == 0 else "text"
                    ]


def _chart(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    categories = spec.get("categories") or []
    series = spec.get("series") or []
    if (
        not isinstance(categories, list)
        or not categories
        or not isinstance(series, list)
    ):
        _bullets(slide, spec.get("bullets") or [], theme)
        return
    data = ChartData()
    data.categories = [str(x) for x in categories]
    series_count = 0
    for item in series:
        if isinstance(item, dict) and isinstance(item.get("values"), list):
            values = [float(x) for x in item["values"]]
            if len(values) == len(data.categories):
                data.add_series(str(item.get("name") or "系列"), values)
                series_count += 1
    if not series_count:
        _bullets(slide, spec.get("bullets") or [], theme)
        return
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.8),
        Inches(1.7),
        Inches(11.5),
        Inches(4.9),
        data,
    ).chart
    chart.has_legend = len(series) > 1
    chart.value_axis.has_major_gridlines = True
    for item in chart.series:
        item.format.fill.solid()
        item.format.fill.fore_color.rgb = theme["colors"]["primary"]
    for axis in (chart.category_axis, chart.value_axis):
        axis.tick_labels.font.name = theme["fonts"]["body"]
        axis.tick_labels.font.color.rgb = theme["colors"]["text"]
    if chart.has_legend:
        chart.legend.font.name = theme["fonts"]["body"]


def _timeline(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    events = spec["events"]
    _title(slide, str(spec.get("title") or ""), theme, str(spec.get("subtitle") or ""))
    left, width, y = 1.0, 11.3, 3.55
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(y), Inches(width), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = theme["colors"]["primary"]
    line.line.fill.background()
    step = width / max(1, len(events) - 1)
    box_width = min(2.55, 10.8 / len(events))
    for index, event in enumerate(events):
        x = left + index * step
        marker = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x - 0.16),
            Inches(y - 0.14),
            Inches(0.34),
            Inches(0.34),
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = theme["colors"]["primary"]
        marker.line.fill.background()
        box_x = min(12.55 - box_width, max(0.55, x - box_width / 2))
        box = slide.shapes.add_textbox(
            Inches(box_x),
            Inches(2.02 if index % 2 == 0 else 3.88),
            Inches(box_width),
            Inches(1.32),
        )
        box.text_frame.text = f"{event.get('label', '')}\n{event.get('title', '')}\n{event.get('detail', '')}"
        _text(box, theme, 18)
        for paragraph in box.text_frame.paragraphs:
            if paragraph.runs:
                paragraph.runs[0].font.bold = True
                break


def _icon_grid(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    items = spec["items"]
    _title(slide, str(spec.get("title") or ""), theme, str(spec.get("subtitle") or ""))
    columns = 4 if len(items) > 4 else 2 if len(items) > 2 else len(items)
    rows = math.ceil(len(items) / columns)
    card_width = 11.8 / columns - 0.25
    card_height = 4.65 / rows - 0.25
    for index, item in enumerate(items):
        column, row = index % columns, index // columns
        x = 0.65 + column * (12.0 / columns)
        y = 1.85 + row * (4.65 / rows)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(card_width),
            Inches(card_height),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = theme["colors"]["pale"]
        card.line.color.rgb = theme["colors"]["primary"]
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x + 0.2),
            Inches(y + 0.28),
            Inches(0.78),
            Inches(0.78),
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = theme["colors"]["primary"]
        badge.line.fill.background()
        badge.text_frame.text = _ICON_LABELS[item["icon"]]
        _text(badge, theme, 9, "inverse", True)
        badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        box = slide.shapes.add_textbox(
            Inches(x + 1.15),
            Inches(y + 0.25),
            Inches(card_width - 1.45),
            Inches(card_height - 0.4),
        )
        box.text_frame.text = f"{item.get('title', '')}\n{item.get('detail', '')}"
        _text(box, theme, 14)
        box.text_frame.paragraphs[0].runs[0].font.bold = True


def _route_map(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    _title(slide, str(spec.get("title") or ""), theme, str(spec.get("subtitle") or ""))
    sea = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(6.25), Inches(1.75), Inches(0.8), Inches(4.8)
    )
    sea.fill.solid()
    sea.fill.fore_color.rgb = RGBColor(0xB9, 0xDB, 0xF4)
    sea.line.fill.background()
    for text, x in (("EUROPE", 1.0), ("ASIA", 9.6)):
        label = slide.shapes.add_textbox(
            Inches(x), Inches(2.0), Inches(2.2), Inches(0.5)
        )
        label.text_frame.text = text
        _text(label, theme, 20, "muted", True)
    points = spec["points"]
    positions: dict[str, int] = {"europe": 0, "route": 0, "asia": 0}
    totals = {
        side: sum(point["side"] == side for point in points) for side in positions
    }
    for point in points:
        side = point["side"]
        x = 2.0 if side == "europe" else 9.65 if side == "asia" else 6.48
        slot = positions[side]
        positions[side] += 1
        y = 2.75 + slot * (2.9 / max(1, totals[side] - 1))
        marker = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.34), Inches(0.34)
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = theme["colors"]["primary"]
        marker.line.fill.background()
        box_x = x + 0.48
        box = slide.shapes.add_textbox(
            Inches(box_x), Inches(y - 0.08), Inches(2.25), Inches(0.9)
        )
        box.text_frame.text = f"{point.get('name', '')}\n{point.get('detail', '')}"
        _text(box, theme, 12)


def _add_picture_fit(
    slide, raw: bytes, *, x: float, y: float, w: float, h: float, name: str
):
    with Image.open(BytesIO(raw)) as image:
        ratio = image.width / image.height
    frame_ratio = w / h
    if ratio >= frame_ratio:
        picture_w, picture_h = w, w / ratio
        picture_x, picture_y = x, y + (h - picture_h) / 2
    else:
        picture_h, picture_w = h, h * ratio
        picture_x, picture_y = x + (w - picture_w) / 2, y
    picture = slide.shapes.add_picture(
        BytesIO(raw),
        Inches(picture_x),
        Inches(picture_y),
        width=Inches(picture_w),
        height=Inches(picture_h),
    )
    picture.name = name
    return picture


def _hero_photo(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    raw, _ = _validated_material(spec["photo"], expected={"image/png", "image/jpeg"})
    _add_picture_fit(slide, raw, x=0, y=0, w=13.333, h=7.5, name="photo:cover")
    panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(4.55), Inches(7.2), Inches(2.15)
    )
    panel.name = "panel:hero"
    panel.fill.solid()
    panel.fill.fore_color.rgb = theme["colors"]["text"]
    panel.fill.transparency = 10
    panel.line.fill.background()
    title = slide.shapes.add_textbox(
        Inches(1.0), Inches(4.82), Inches(6.5), Inches(0.9)
    )
    title.name = "text:hero-title"
    title.text_frame.text = str(spec.get("title") or "")
    _text(title, theme, 34, "inverse", True, "title")
    subtitle = slide.shapes.add_textbox(
        Inches(1.02), Inches(5.78), Inches(6.3), Inches(0.58)
    )
    subtitle.name = "text:hero-subtitle"
    subtitle.text_frame.text = str(spec.get("subtitle") or "")
    _text(subtitle, theme, 18, "inverse")


def _photo_collage(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    _title(slide, str(spec.get("title") or ""), theme, str(spec.get("subtitle") or ""))
    photos = spec["photos"]
    columns = 2
    rows = math.ceil(len(photos) / columns)
    frame_w, frame_h = 5.75, 4.45 / rows
    for index, photo in enumerate(photos):
        raw, _ = _validated_material(
            photo["material"], expected={"image/png", "image/jpeg"}
        )
        col, row = index % columns, index // columns
        x, y = 0.7 + col * 6.05, 1.85 + row * (frame_h + 0.25)
        backdrop = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(frame_w), Inches(frame_h)
        )
        backdrop.name = f"photo-frame:{index + 1}"
        backdrop.fill.solid()
        backdrop.fill.fore_color.rgb = theme["colors"]["pale"]
        backdrop.line.fill.background()
        _add_picture_fit(
            slide,
            raw,
            x=x,
            y=y,
            w=frame_w,
            h=frame_h - 0.42,
            name=f"photo:scene-{index + 1}",
        )
        caption = slide.shapes.add_textbox(
            Inches(x + 0.1),
            Inches(y + frame_h - 0.4),
            Inches(frame_w - 0.2),
            Inches(0.3),
        )
        caption.name = f"source-caption:{index + 1}"
        caption.text_frame.text = str(photo.get("caption") or "")
        _text(caption, theme, 10, "muted")


def _native_svg(
    slide,
    raw: bytes,
    theme: dict[str, Any],
    *,
    x: float,
    y: float,
    w: float,
    h: float,
    prefix: str,
) -> int:
    view, primitives = parse_editable_svg(raw)
    vx, vy, vw, vh = view
    sx, sy = w / vw, h / vh
    count = 0
    for primitive in primitives:
        values = primitive.values
        fill = (
            RGBColor.from_string((primitive.fill or "#000000")[1:])
            if primitive.fill
            else None
        )
        stroke = (
            RGBColor.from_string((primitive.stroke or "#000000")[1:])
            if primitive.stroke
            else theme["colors"]["primary"]
        )
        if primitive.kind == "circle":
            cx, cy, radius = values
            shape = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(x + (cx - radius - vx) * sx),
                Inches(y + (cy - radius - vy) * sy),
                Inches(2 * radius * sx),
                Inches(2 * radius * sy),
            )
            if fill:
                shape.fill.solid()
                shape.fill.fore_color.rgb = fill
            else:
                shape.fill.background()
            shape.line.color.rgb = stroke
            shape.name = f"vector:{prefix}:{count + 1}"
            count += 1
        elif primitive.kind == "rect":
            px, py, pw, ph = values
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x + (px - vx) * sx),
                Inches(y + (py - vy) * sy),
                Inches(pw * sx),
                Inches(ph * sy),
            )
            if fill:
                shape.fill.solid()
                shape.fill.fore_color.rgb = fill
            else:
                shape.fill.background()
            shape.line.color.rgb = stroke
            shape.name = f"vector:{prefix}:{count + 1}"
            count += 1
        else:
            coords = values
            for offset in range(0, len(coords) - 2, 2):
                x1, y1, x2, y2 = coords[offset : offset + 4]
                line = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT,
                    Inches(x + (x1 - vx) * sx),
                    Inches(y + (y1 - vy) * sy),
                    Inches(x + (x2 - vx) * sx),
                    Inches(y + (y2 - vy) * sy),
                )
                line.line.color.rgb = stroke
                line.line.width = Pt(max(1, primitive.stroke_width * min(sx, sy) * 12))
                line.name = f"vector:{prefix}:{count + 1}"
                count += 1
    return count


def _icon_facts(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    _title(slide, str(spec.get("title") or ""), theme, str(spec.get("subtitle") or ""))
    items = spec["items"]
    columns = 3 if len(items) > 4 else 2
    rows = math.ceil(len(items) / columns)
    card_w, card_h = 11.85 / columns - 0.22, 4.7 / rows - 0.22
    for index, item in enumerate(items):
        col, row = index % columns, index // columns
        x, y = 0.68 + col * (12.0 / columns), 1.78 + row * (4.7 / rows)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(card_w),
            Inches(card_h),
        )
        card.name = f"card:icon-{index + 1}"
        card.fill.solid()
        card.fill.fore_color.rgb = theme["colors"]["pale"]
        card.line.fill.background()
        raw, _ = _validated_material(item["icon"], expected={"image/svg+xml"})
        _native_svg(
            slide,
            raw,
            theme,
            x=x + 0.18,
            y=y + 0.22,
            w=0.72,
            h=0.72,
            prefix=str(index + 1),
        )
        box = slide.shapes.add_textbox(
            Inches(x + 1.08),
            Inches(y + 0.2),
            Inches(card_w - 1.25),
            Inches(card_h - 0.35),
        )
        box.name = f"text:icon-{index + 1}"
        box.text_frame.text = f"{item['title']}\n{item.get('detail', '')}"
        _text(box, theme, 18)
        box.text_frame.paragraphs[0].runs[0].font.bold = True


def _geo_route_map(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    _title(slide, str(spec.get("title") or ""), theme, str(spec.get("subtitle") or ""))
    points = validate_geo_points(spec["points"])
    x, y, w, h = 0.7, 1.75, 11.9, 4.9
    map_material = spec.get("map")
    if map_material is not None:
        raw, _ = _validated_material(
            map_material, expected={"image/png", "image/jpeg"}
        )
        _add_picture_fit(
            slide, raw, x=x, y=y, w=w, h=h, name="map:base-real"
        )
    else:
        sea = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        sea.name = "map:bosphorus-water"
        sea.fill.solid()
        sea.fill.fore_color.rgb = RGBColor(0xA9, 0xD8, 0xE8)
        sea.line.fill.background()
        west = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w * 0.54), Inches(h)
        )
        west.name = "map:europe-land"
        west.fill.solid()
        west.fill.fore_color.rgb = RGBColor(0xE8, 0xDF, 0xC8)
        west.line.fill.background()
        east = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x + w * 0.62),
            Inches(y),
            Inches(w * 0.38),
            Inches(h),
        )
        east.name = "map:asia-land"
        east.fill.solid()
        east.fill.fore_color.rgb = RGBColor(0xD7, 0xD7, 0xBD)
        east.line.fill.background()
    for label, lx in (
        ("EUROPE", x + 0.35),
        ("ASIA", x + 9.55),
        ("BOSPHORUS", x + 6.45),
    ):
        box = slide.shapes.add_textbox(
            Inches(lx), Inches(y + 0.2), Inches(1.9), Inches(0.35)
        )
        box.name = f"map:label-{label.lower()}"
        box.text_frame.text = label
        _text(box, theme, 18, "muted", True)
    if spec.get("attribution"):
        attribution = slide.shapes.add_textbox(
            Inches(x + 3.75), Inches(y + h - 0.28), Inches(4.4), Inches(0.22)
        )
        attribution.name = "map:attribution"
        attribution.text_frame.text = str(spec["attribution"])
        _text(attribution, theme, 8, "muted")
    projected = []
    for point in points:
        px, py = project_point(point, width=w, height=h)
        projected.append((x + px, y + py))
    for index in range(len(projected) - 1):
        a, b = projected[index], projected[index + 1]
        route = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(a[0]),
            Inches(a[1]),
            Inches(b[0]),
            Inches(b[1]),
        )
        route.name = f"map:route-{index + 1}"
        route.line.color.rgb = theme["colors"]["primary"]
        route.line.width = Pt(4)
    side_slots = {"europe": 0, "asia": 0}
    side_totals = {
        side: sum(point.side == side for point in points) for side in side_slots
    }
    for index, (point, (px, py)) in enumerate(zip(points, projected), 1):
        marker = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(px - 0.11),
            Inches(py - 0.11),
            Inches(0.22),
            Inches(0.22),
        )
        marker.name = f"map:landmark-{index}"
        marker.fill.solid()
        marker.fill.fore_color.rgb = theme["colors"]["primary"]
        marker.line.color.rgb = theme["colors"]["inverse"]
        slot = side_slots[point.side]
        side_slots[point.side] += 1
        label_x = x + 0.28 if point.side == "europe" else x + 9.35
        usable = 3.45
        label_y = y + 0.82 + (usable * slot / max(1, side_totals[point.side] - 1))
        anchor_x = label_x + 2.05 if point.side == "europe" else label_x
        callout = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(px),
            Inches(py),
            Inches(anchor_x),
            Inches(label_y + 0.18),
        )
        callout.name = f"map:callout-{index}"
        callout.line.color.rgb = theme["colors"]["muted"]
        callout.line.width = Pt(1)
        box = slide.shapes.add_textbox(
            Inches(label_x), Inches(label_y), Inches(2.15), Inches(0.42)
        )
        box.name = f"map:landmark-label-{index}"
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        box.fill.transparency = 10
        box.line.fill.background()
        box.text_frame.text = f"{index}. {point.name}"
        _text(box, theme, 18, "text", True)


def _data_story(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    _title(slide, str(spec.get("title") or ""), theme, str(spec.get("subtitle") or ""))
    metric = slide.shapes.add_textbox(
        Inches(0.85), Inches(2.0), Inches(4.0), Inches(1.35)
    )
    metric.name = "text:metric"
    metric.text_frame.text = f"{spec.get('metric', '')}{spec.get('unit', '')}"
    _text(metric, theme, 48, "primary", True, "title")
    body = slide.shapes.add_textbox(
        Inches(0.9), Inches(3.45), Inches(4.1), Inches(2.25)
    )
    body.name = "text:data-body"
    body.text_frame.text = str(spec.get("body") or "")
    _text(body, theme, 20)
    _bullets(slide, spec["facts"], theme, x=5.35, y=1.95, w=7.0, h=4.45)


def _quote_photo(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    raw, _ = _validated_material(spec["photo"], expected={"image/png", "image/jpeg"})
    _add_picture_fit(slide, raw, x=0, y=0, w=6.15, h=7.5, name="photo:quote")
    quote = slide.shapes.add_textbox(
        Inches(6.75), Inches(1.5), Inches(5.65), Inches(3.6)
    )
    quote.name = "text:quote"
    quote.text_frame.text = f"“{spec.get('quote', '')}”"
    _text(quote, theme, 28, "text", True, "title")
    author = slide.shapes.add_textbox(
        Inches(6.8), Inches(5.35), Inches(5.2), Inches(0.65)
    )
    author.name = "text:attribution"
    author.text_frame.text = str(spec.get("attribution") or "")
    _text(author, theme, 18, "muted")


def _image(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    _title(slide, str(spec.get("title") or ""), theme, str(spec.get("subtitle") or ""))
    raw = base64.b64decode(str(spec["image_data"]).split(",", 1)[1], validate=True)
    with Image.open(BytesIO(raw)) as image:
        image_ratio = image.width / image.height
    frame_ratio = 11.85 / 4.75
    picture = slide.shapes.add_picture(
        BytesIO(raw),
        Inches(0.75),
        Inches(1.75),
        width=Inches(11.85),
        height=Inches(4.75),
    )
    if image_ratio > frame_ratio:
        crop = (1 - frame_ratio / image_ratio) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    elif image_ratio < frame_ratio:
        crop = (1 - image_ratio / frame_ratio) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    if spec.get("caption"):
        caption = slide.shapes.add_textbox(
            Inches(0.8), Inches(6.55), Inches(11.5), Inches(0.3)
        )
        caption.text_frame.text = str(spec["caption"])
        _text(caption, theme, 10, "muted")


def build_pptx(content: str) -> bytes:
    value = _spec(content)
    theme = _theme(value)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for index, item in enumerate(value["slides"]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = theme["colors"]["background"]
        kind = str(item.get("layout") or "bullets")
        title = str(item.get("title") or f"第 {index + 1} 页")
        if kind == "hero_photo":
            _hero_photo(slide, item, theme)
        elif kind == "photo_collage":
            _photo_collage(slide, item, theme)
        elif kind == "geo_route_map":
            _geo_route_map(slide, item, theme)
        elif kind == "icon_facts":
            _icon_facts(slide, item, theme)
        elif kind == "quote_photo":
            _quote_photo(slide, item, theme)
        elif kind == "data_story":
            _data_story(slide, item, theme)
        elif kind == "title":
            box = slide.shapes.add_textbox(
                Inches(1), Inches(2.1), Inches(11.3), Inches(1.3)
            )
            box.text_frame.text = title
            _text(box, theme, 36, bold=True, font="title")
            box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            sub = slide.shapes.add_textbox(
                Inches(1.3), Inches(3.65), Inches(10.7), Inches(0.8)
            )
            sub.text_frame.text = str(
                item.get("subtitle") or value.get("subtitle") or ""
            )
            _text(sub, theme, 18, "muted")
            sub.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            accent = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0),
                Inches(0),
                Inches(0.28),
                Inches(7.5),
            )
            accent.fill.solid()
            accent.fill.fore_color.rgb = theme["colors"]["primary"]
            accent.line.fill.background()
            halo = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(10.65),
                Inches(0.55),
                Inches(1.75),
                Inches(1.75),
            )
            halo.fill.solid()
            halo.fill.fore_color.rgb = theme["colors"]["pale"]
            halo.line.fill.background()
        elif kind == "section":
            background.fore_color.rgb = theme["colors"]["primary"]
            box = slide.shapes.add_textbox(
                Inches(1), Inches(2.4), Inches(11.3), Inches(1.4)
            )
            box.text_frame.text = title
            _text(box, theme, 36, "inverse", True, "title")
        elif kind == "two_column":
            _title(slide, title, theme, str(item.get("subtitle") or ""))
            _bullets(slide, item.get("left") or [], theme, x=0.7, w=5.7)
            _bullets(slide, item.get("right") or [], theme, x=6.9, w=5.7)
        elif kind == "chart":
            _title(slide, title, theme)
            _chart(slide, item, theme)
        elif kind == "table":
            _title(slide, title, theme)
            _table(slide, item, theme)
        elif kind == "timeline":
            _timeline(slide, item, theme)
        elif kind == "icon_grid":
            _icon_grid(slide, item, theme)
        elif kind == "route_map":
            _route_map(slide, item, theme)
        elif kind == "image":
            _image(slide, item, theme)
        else:
            _title(slide, title, theme, str(item.get("subtitle") or ""))
            _bullets(slide, item.get("bullets") or [], theme)
        marker = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.01), Inches(0.01)
        )
        marker.name = f"layout:{kind}"
        marker.fill.background()
        marker.line.fill.background()
        number = slide.shapes.add_textbox(
            Inches(12.25), Inches(7.02), Inches(0.45), Inches(0.25)
        )
        number.name = "page-number:footer"
        number.text_frame.text = str(index + 1)
        _text(number, theme, 9, "muted")
    buffer = BytesIO()
    prs.save(buffer)
    data = buffer.getvalue()
    Presentation(BytesIO(data))
    return data


def render_pptx_pdf(pptx_path: Path) -> bytes:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise RuntimeError("PPTX 预览渲染器不可用（未检测到 LibreOffice）")
    with tempfile.TemporaryDirectory(prefix="ppt-render-") as output:
        profile = (Path(output) / "profile").as_uri()
        result = subprocess.run(
            [
                executable,
                f"-env:UserInstallation={profile}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                output,
                str(pptx_path),
            ],
            capture_output=True,
            text=True,
            timeout=90,
        )
        path = Path(output) / f"{pptx_path.stem}.pdf"
        if result.returncode or not path.is_file():
            raise RuntimeError(
                f"PPTX 预览渲染失败（exit {result.returncode}）：{(result.stderr or result.stdout).strip()[:240]}"
            )
        data = path.read_bytes()
        if not data.startswith(b"%PDF-"):
            raise RuntimeError("PPTX 预览渲染器未生成有效 PDF")
        return data
