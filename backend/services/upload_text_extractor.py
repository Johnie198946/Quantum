"""Extract bounded text from contribution uploads before Hermes sees them."""
from __future__ import annotations

import io
from pathlib import Path
from typing import Callable
import hashlib
import zipfile

MAX_EXTRACTED_CHARACTERS = 200_000


def _bounded(text: str) -> str:
    value = text.strip()
    if not value:
        raise ValueError("uploaded file contains no extractable text")
    if len(value) > MAX_EXTRACTED_CHARACTERS:
        raise ValueError("uploaded file extracted text exceeds contribution limit")
    return value


def extract_uploaded_text(data: bytes, *, filename: str, content_type: str, analyze_images: Callable[[list[bytes]], list[str]] | None = None) -> str:
    suffix = Path(filename).suffix.lower()
    mime = content_type.split(";", 1)[0].strip().lower()
    if mime.startswith("text/") or suffix in {".md", ".txt", ".csv", ".json", ".yaml", ".yml", ".html"}:
        return _bounded(data.decode("utf-8-sig"))
    if suffix in {".doc", ".ppt"}:
        import tempfile
        from backend.services.presentation_renderer import render_office_pdf
        with tempfile.TemporaryDirectory(prefix="office-source-") as directory:
            original = Path(directory) / ("source" + suffix)
            original.write_bytes(data)
            converted = render_office_pdf(original)
        return extract_uploaded_text(converted, filename="source.pdf", content_type="application/pdf", analyze_images=analyze_images)
    image_parts = []
    if analyze_images is not None and suffix in {".docx", ".pptx", ".pdf"}:
        images = {}
        if suffix in {".docx", ".pptx"}:
            prefix = "word/media/" if suffix == ".docx" else "ppt/media/"
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                total = 0
                for item in archive.infolist():
                    if item.is_dir() or not item.filename.startswith(prefix):
                        continue
                    total += item.file_size
                    if total > 25 * 1024 * 1024:
                        raise ValueError("embedded images exceed 25 MB extraction limit")
                    image = archive.read(item)
                    images.setdefault(hashlib.sha256(image).hexdigest(), image)
        else:
            from pypdf import PdfReader
            total = 0
            for page in PdfReader(io.BytesIO(data)).pages:
                for item in page.images:
                    image = item.data
                    total += len(image)
                    if total > 25 * 1024 * 1024:
                        raise ValueError("embedded images exceed 25 MB extraction limit")
                    images.setdefault(hashlib.sha256(image).hexdigest(), image)
        if len(images) > 16:
            raise ValueError("more than 16 distinct embedded images; split the document before analysis")
        if images:
            analyses = analyze_images(list(images.values()))
            if len(analyses) != len(images) or any(not text.strip() for text in analyses):
                raise ValueError("embedded image analysis incomplete")
            image_parts = [f"[内嵌图片 {index + 1}，SHA256 {digest}；模型识别，需核对原件]\n{text}"
                           for index, (digest, text) in enumerate(zip(images, analyses))]
    def finish(text):
        return _bounded("\n\n".join([text, *image_parts]))
    if suffix == ".pdf" or mime == "application/pdf":
        from pypdf import PdfReader
        return finish("\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages))
    if suffix == ".docx" or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        from docx import Document
        document = Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        parts.extend("\t".join(cell.text for cell in row.cells)
                     for table in document.tables for row in table.rows)
        return finish("\n".join(parts))
    if suffix == ".xlsx" or mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"# {sheet.title}")
            lines.extend("\t".join("" if value is None else str(value) for value in row)
                         for row in sheet.iter_rows(values_only=True))
        return _bounded("\n".join(lines))
    if suffix == ".pptx" or mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(data))
        return finish("\n".join(
            shape.text for slide in presentation.slides for shape in slide.shapes
            if hasattr(shape, "text") and shape.text
        ))
    raise ValueError("uploaded file type has no safe text extractor")
