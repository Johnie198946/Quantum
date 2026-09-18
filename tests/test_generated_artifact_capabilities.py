from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

from backend.capability_handlers import HANDLERS
from backend.services.generated_artifacts import (
    GeneratedArtifactError,
    analyze_data,
    create_media,
    create_pdf,
    create_spreadsheet,
    generated_artifact_path,
    read_generated_artifact,
)


def test_generated_artifacts_are_real_hashed_and_owner_bound(tmp_path, monkeypatch):
    monkeypatch.setenv("AI_LAB_GENERATED_ARTIFACT_ROOT", str(tmp_path))
    sheet = create_spreadsheet(
        "tenant-a", "user-a",
        {"title": "sales", "columns": ["month", "value"], "rows": [["Jan", 3], ["Feb", 5]]},
    )
    sheet_path, sheet_receipt = generated_artifact_path("tenant-a", "user-a", sheet["artifact_id"])
    assert sheet_path.read_bytes().startswith(b"PK")
    assert load_workbook(sheet_path).active.max_row == 3
    assert hashlib.sha256(sheet_path.read_bytes()).hexdigest() == sheet_receipt["content_hash"]

    pdf = create_pdf("tenant-a", "user-a", {"title": "report", "content": "verified report\nline two"})
    pdf_path, _ = generated_artifact_path("tenant-a", "user-a", pdf["artifact_id"])
    assert pdf_path.read_bytes().startswith(b"%PDF-")

    with pytest.raises(GeneratedArtifactError, match="not found"):
        read_generated_artifact("tenant-a", "user-b", sheet["artifact_id"])

    sheet_path.write_bytes(b"corrupt")
    with pytest.raises(GeneratedArtifactError, match="hash mismatch"):
        generated_artifact_path("tenant-a", "user-a", sheet["artifact_id"])


def test_data_analysis_and_media_create_real_artifacts(tmp_path, monkeypatch):
    monkeypatch.setenv("AI_LAB_GENERATED_ARTIFACT_ROOT", str(tmp_path))
    analysis = analyze_data(
        "tenant-a", "user-a",
        {"columns": ["value", "label"], "rows": [[1, "a"], [3, "b"], [None, "c"]]},
    )
    assert analysis["analysis"]["columns"]["value"] == {
        "missing": 1, "count": 2, "min": 1.0, "max": 3.0, "mean": 2.0,
    }
    analysis_path, _ = generated_artifact_path("tenant-a", "user-a", analysis["artifact_id"])
    assert json.loads(analysis_path.read_text())["row_count"] == 3

    media = create_media("tenant-a", "user-a", {"prompt": "Quarterly result", "width": 640, "height": 360})
    media_path, receipt = generated_artifact_path("tenant-a", "user-a", media["artifact_id"])
    with Image.open(media_path) as image:
        assert image.format == "PNG"
        assert image.size == (640, 360)
    assert receipt["metadata"]["prompt_hash"] != "Quarterly result"


def test_generated_artifact_handlers_are_registered():
    assert {
        "office.spreadsheet.create", "office.pdf.create", "data.analyze", "media.create"
    }.issubset(HANDLERS)
