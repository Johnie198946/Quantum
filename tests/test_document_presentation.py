import hashlib
import asyncio
import base64
import copy
import json
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
import httpx
from fastapi import FastAPI
from docx import Document
from PIL import Image
from pptx import Presentation
from pypdf import PdfWriter

from backend.services import document_sources
from backend.services.document_sources import DocumentSourceError
from backend.services.presentation_renderer import build_pptx, render_pptx_pdf
from backend.services.presentation_scenario import build_presentation_plan
from backend.services.dsl_safety_compiler import DSLSafetyCompiler
from backend.services.workflow_executor import trusted_task_agent_config
from backend.api.auth import require_auth
from backend.api.documents import router as documents_router


THEME_A = {
    "colors": {
        "primary": "#1A73E8",
        "text": "#102030",
        "muted": "#405060",
        "pale": "#DDEEFF",
        "background": "#FAFBFC",
        "inverse": "#FFFFFF",
    },
    "fonts": {"title": "Arial", "body": "Courier New"},
}
THEME_B = {
    "colors": {
        "primary": "#D93025",
        "text": "#302010",
        "muted": "#605040",
        "pale": "#FFEEDD",
        "background": "#FFFDFC",
        "inverse": "#FFFF00",
    },
    "fonts": {"title": "Georgia", "body": "Times New Roman"},
}


def _docx(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _request(app: FastAPI, method: str, path: str, **kwargs):
    async def run():
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://testserver"
        ) as client:
            return await client.request(method, path, **kwargs)

    return asyncio.run(run())


def test_private_document_is_tenant_bound_and_original_survives_parse_failure(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )
    data = _docx("季度收入增长，客户留存改善。")
    receipt = document_sources.save_document_source(
        tenant_key="tenant-a",
        user_id="user-a",
        filename="report.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=data,
        expected_hash=hashlib.sha256(data).hexdigest(),
    )
    assert receipt["status"] == "ready"
    assert receipt["note_id"] == receipt["source_id"]
    note_path = (
        document_sources.note_directory("tenant-a", "user-a")
        / f"{receipt['source_id']}.md"
    )
    assert note_path.is_file()
    assert "季度收入增长" in note_path.read_text(encoding="utf-8")
    assert document_sources.document_text("tenant-a", "user-a", receipt["source_id"])[
        0
    ].startswith("季度收入")
    extracted = (
        document_sources.note_directory("tenant-a", "user-a")
        / ".documents"
        / receipt["source_id"]
        / "extracted.txt"
    )
    extracted.write_text("被篡改的提取内容", encoding="utf-8")
    with pytest.raises(DocumentSourceError, match="完整性") as integrity:
        document_sources.document_text("tenant-a", "user-a", receipt["source_id"])
    assert integrity.value.code == "document_integrity_error"
    with pytest.raises(DocumentSourceError, match="文档不存在"):
        document_sources.read_document_receipt(
            "tenant-a", "user-b", receipt["source_id"]
        )

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    output = BytesIO()
    writer.write(output)
    failed = document_sources.save_document_source(
        tenant_key="tenant-a",
        user_id="user-a",
        filename="scan.pdf",
        content_type="application/pdf",
        data=output.getvalue(),
    )
    assert failed["status"] == "parse_failed"
    original, _ = document_sources.document_original_path(
        "tenant-a", "user-a", failed["source_id"]
    )
    assert original.read_bytes() == output.getvalue()
    assert failed["parse_error"]["code"] == "no_extractable_text"
    encrypted_writer = PdfWriter()
    encrypted_writer.add_blank_page(width=100, height=100)
    encrypted_writer.encrypt("secret")
    encrypted_output = BytesIO()
    encrypted_writer.write(encrypted_output)
    encrypted = document_sources.save_document_source(
        tenant_key="tenant-a",
        user_id="user-a",
        filename="locked.pdf",
        content_type="application/pdf",
        data=encrypted_output.getvalue(),
    )
    assert (
        encrypted["status"] == "parse_failed"
        and encrypted["parse_error"]["code"] == "encrypted_pdf"
    )


def test_task_agent_manifest_is_projected_to_strict_bridge_schema():
    agent = type("Agent", (), {
        "id": "agent-private",
        "private_prompt_delta": "approved prompt",
        "composition_manifest": {
            "capability_agent_ids": ["main_agent"],
            "invoked_agent_ids": ["tenant_specialist"],
            "delegation": {"max_concurrent_children": 3, "max_spawn_depth": 1},
            "knowledge_scope": ["knowledge/product/public"],
            "plan_id": "must-not-cross-runtime-boundary",
        },
    })()
    assert trusted_task_agent_config(agent) == {  # type: ignore[arg-type]
        "id": "agent-private",
        "prompt": "approved prompt",
        "capability_agent_ids": ["main_agent", "tenant_specialist"],
        "knowledge_scope": ["knowledge/product/public"],
        "delegation": {"max_concurrent_children": 3, "max_spawn_depth": 1},
    }


def test_document_rejects_legacy_and_oversize(tmp_path, monkeypatch):
    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )
    with pytest.raises(DocumentSourceError) as legacy:
        document_sources.save_document_source(
            tenant_key="t",
            user_id="u",
            filename="old.doc",
            content_type="application/msword",
            data=b"x",
        )
    assert legacy.value.code == "legacy_doc_unsupported"
    with pytest.raises(DocumentSourceError) as oversized:
        document_sources.save_document_source(
            tenant_key="t",
            user_id="u",
            filename="large.pdf",
            content_type="application/pdf",
            data=b"x" * (document_sources.MAX_DOCUMENT_BYTES + 1),
        )
    assert oversized.value.code == "document_too_large"


def test_authenticated_upload_receipt_download_hash_and_user_boundary(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )
    app = FastAPI()
    app.include_router(documents_router)
    identity = {"tenant_key": "tenant-a", "user_id": "user-a"}
    app.dependency_overrides[require_auth] = lambda: identity
    data = _docx("可验证的私有上传")
    digest = hashlib.sha256(data).hexdigest()
    response = _request(
        app,
        "POST",
        "/api/v1/documents",
        content=data,
        headers={
            "X-File-Name": "private.docx",
            "X-Content-Hash": digest,
            "X-File-Opt-Out": "true",
            "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    )
    assert response.status_code == 201
    receipt = response.json()
    assert (
        receipt["content_hash"] == digest
        and receipt["contribution_status"] == "excluded"
    )
    downloaded = _request(
        app, "GET", f"/api/v1/documents/{receipt['source_id']}/download"
    )
    assert (
        downloaded.content == data and downloaded.headers["x-content-sha256"] == digest
    )
    identity["user_id"] = "user-b"
    assert (
        _request(app, "GET", f"/api/v1/documents/{receipt['source_id']}").status_code
        == 404
    )


def test_authorized_default_upload_uses_authenticated_identity_and_real_queue_receipt(
    tmp_path, monkeypatch
):
    from agreement_fixtures import set_user_contribution_consent
    from backend.db import SessionLocal
    from backend.models.knowledge_contribution import KnowledgeContributionOutbox
    from backend.services.knowledge_contribution import set_contribution_policy

    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )
    monkeypatch.setenv("HERMES_CHAT_RUN_DB", str(tmp_path / "runs.sqlite3"))
    tenant, user = "document-" + uuid4().hex, "user-" + uuid4().hex

    async def authorize():
        await set_contribution_policy(
            tenant_key=tenant,
            enabled=True,
            agreement_version="v4",
            effective_at=datetime.now(timezone.utc) - timedelta(minutes=1),
        )
        await set_user_contribution_consent(
            tenant_key=tenant,
            user_id=user,
            service_agreement_version="service-2026-09-06",
            participation_enabled=True,
        )

    asyncio.run(authorize())
    app = FastAPI()
    app.include_router(documents_router)
    app.dependency_overrides[require_auth] = lambda: {"tenant_key": tenant, "sub": user}
    response = _request(
        app,
        "POST",
        "/api/v1/documents",
        content=_docx("参与贡献但原件保持私有"),
        headers={
            "X-File-Name": "governed.docx",
            "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    )
    receipt = response.json()
    assert response.status_code == 201 and receipt["contribution_status"] == "queued"
    current = _request(app, "GET", f"/api/v1/documents/{receipt['source_id']}")
    assert current.status_code == 200
    assert current.json()["contribution_status"] == "compiling"
    assert current.json()["note_id"] == receipt["source_id"]

    async def load_event():
        async with SessionLocal() as db:
            return await db.get(
                KnowledgeContributionOutbox, receipt["contribution_event_id"]
            )

    event = asyncio.run(load_event())
    assert event and (event.tenant_key, event.user_id, event.source_id) == (
        tenant,
        user,
        receipt["source_id"],
    )
    assert receipt["contribution_run_id"]


def test_default_upload_without_server_authorization_is_denied_but_remains_private(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )
    app = FastAPI()
    app.include_router(documents_router)
    identity = {
        "tenant_key": "unauthorized-" + uuid4().hex,
        "user_id": "user-" + uuid4().hex,
    }
    app.dependency_overrides[require_auth] = lambda: identity
    receipt = _request(
        app,
        "POST",
        "/api/v1/documents",
        content=_docx("未授权不会贡献"),
        headers={
            "X-File-Name": "denied.docx",
            "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    ).json()
    assert receipt["contribution_status"] == "denied"
    assert (
        _request(
            app, "GET", f"/api/v1/documents/{receipt['source_id']}/download"
        ).status_code
        == 200
    )


@pytest.mark.parametrize(
    "result,status",
    [
        (None, "denied"),
        ({}, "failed"),
        (
            {
                "event_id": "contrib-pending",
                "status": "pending",
                "schedule_status": "pending",
                "schedule_error": "queue unavailable",
            },
            "pending",
        ),
    ],
)
def test_upload_never_reports_queued_without_accepted_scheduled_receipt(
    tmp_path, monkeypatch, result, status
):
    import backend.api.documents as api

    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )

    async def rejected(candidate, *, source_content):
        return result

    monkeypatch.setattr(api, "enqueue_and_schedule", rejected)
    app = FastAPI()
    app.include_router(documents_router)
    app.dependency_overrides[require_auth] = lambda: {
        "tenant_key": "tenant-no-auth",
        "user_id": "user-no-auth",
    }
    receipt = _request(
        app,
        "POST",
        "/api/v1/documents",
        content=_docx("私有保存独立成功"),
        headers={
            "X-File-Name": "private.docx",
            "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    ).json()
    assert (
        _request(
            app, "GET", f"/api/v1/documents/{receipt['source_id']}/download"
        ).status_code
        == 200
    )
    assert receipt["status"] == "ready" and receipt["contribution_status"] == status


def test_upload_contribution_failure_does_not_fail_private_save(tmp_path, monkeypatch):
    import backend.api.documents as api

    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )

    async def failed(candidate, *, source_content):
        raise RuntimeError("enqueue failed")

    monkeypatch.setattr(api, "enqueue_and_schedule", failed)
    app = FastAPI()
    app.include_router(documents_router)
    app.dependency_overrides[require_auth] = lambda: {
        "tenant_key": "tenant-fail",
        "user_id": "user-fail",
    }
    receipt = _request(
        app,
        "POST",
        "/api/v1/documents",
        content=_docx("仍保存"),
        headers={
            "X-File-Name": "saved.docx",
            "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    ).json()
    assert (
        _request(
            app, "GET", f"/api/v1/documents/{receipt['source_id']}/download"
        ).status_code
        == 200
    )
    assert (
        receipt["status"] == "ready"
        and receipt["contribution_status"] == "failed"
        and "enqueue failed" in receipt["contribution_error"]
    )


def test_editable_pptx_reopens_and_pdf_preview_uses_same_file(tmp_path):
    spec = {
        "title": "经营复盘",
        "slides": [
            {"layout": "title", "title": "经营复盘", "subtitle": "2026 Q3"},
            {
                "layout": "bullets",
                "title": "关键结论",
                "bullets": ["收入增长", "留存改善"],
            },
            {
                "layout": "two_column",
                "title": "对照",
                "left": ["现状"],
                "right": ["下一步"],
            },
            {
                "layout": "table",
                "title": "指标",
                "headers": ["指标", "值"],
                "rows": [["收入", "100"]],
            },
            {
                "layout": "chart",
                "title": "趋势",
                "categories": ["Q1", "Q2"],
                "series": [{"name": "收入", "values": [80, 100]}],
            },
            {"layout": "conclusion", "title": "结论", "bullets": ["聚焦增长"]},
        ],
    }
    data = build_pptx(json.dumps(spec, ensure_ascii=False))
    assert data.startswith(b"PK")
    reopened = Presentation(BytesIO(data))
    assert len(reopened.slides) == 6
    path = tmp_path / "deck.pptx"
    path.write_bytes(data)
    try:
        pdf = render_pptx_pdf(path)
    except RuntimeError as exc:
        if "预览渲染" in str(exc):
            pytest.skip(str(exc))
        raise
    assert pdf.startswith(b"%PDF-")


def test_visual_presentation_layouts_are_editable_and_images_are_bounded():
    image_bytes = BytesIO()
    Image.new("RGB", (128, 96), "#4A90E2").save(image_bytes, format="PNG")
    image_data = "data:image/png;base64," + base64.b64encode(image_bytes.getvalue()).decode()
    spec = {
        "title": "Istanbul",
        "slides": [
            {
                "layout": "timeline",
                "title": "Three days",
                "events": [
                    {"label": "Day 1", "title": "Old City", "detail": "Hagia Sophia"},
                    {"label": "Day 2", "title": "Bosphorus", "detail": "Ferry"},
                ],
            },
            {
                "layout": "icon_grid",
                "title": "Travel kit",
                "items": [
                    {"icon": "ferry", "title": "Crossing", "detail": "Istanbulkart"},
                    {"icon": "food", "title": "Market", "detail": "Local breakfast"},
                ],
            },
            {
                "layout": "route_map",
                "title": "Europe to Asia",
                "points": [
                    {"name": "Eminönü", "side": "europe", "detail": "Start"},
                    {"name": "Bosphorus", "side": "route", "detail": "Ferry"},
                    {"name": "Kadıköy", "side": "asia", "detail": "Finish"},
                ],
            },
            {
                "layout": "image",
                "title": "Verified image",
                "image_data": image_data,
                "caption": "Embedded upstream asset",
            },
        ],
    }
    reopened = Presentation(BytesIO(build_pptx(json.dumps(spec))))
    assert len(reopened.slides) == 4
    assert all(len(reopened.slides[index].shapes) >= 3 for index in range(3))
    assert any(shape.shape_type == 13 for shape in reopened.slides[3].shapes)
    assert "Old City" in " ".join(
        shape.text for shape in reopened.slides[0].shapes if shape.has_text_frame
    )

    jpeg = BytesIO()
    Image.new("RGB", (128, 96), "#4A90E2").save(jpeg, format="JPEG")
    spec["slides"] = [{
        "layout": "image",
        "title": "Spoofed",
        "image_data": "data:image/png;base64," + base64.b64encode(jpeg.getvalue()).decode(),
    }]
    with pytest.raises(ValueError, match="match MIME type"):
        build_pptx(json.dumps(spec))

    tiny = BytesIO()
    Image.new("RGB", (1, 1), "#4A90E2").save(tiny, format="PNG")
    spec["slides"][0]["image_data"] = (
        "data:image/png;base64," + base64.b64encode(tiny.getvalue()).decode()
    )
    with pytest.raises(ValueError, match="dimensions"):
        build_pptx(json.dumps(spec))


def test_bridge_outline_claim_coverage_rejects_unknown_claims_before_user_approval():
    import scripts.hermes_bridge as bridge

    source_text = "第一条事实。第二条事实。"
    source_id = hashlib.sha256(source_text.encode("utf-8")).hexdigest()
    run = {
        "source_material": {
            "text": source_text,
            "source_id": source_id,
            "source_client_session_id": "ios-session-1",
        }
    }
    with pytest.raises(RuntimeError, match="unknown source claims"):
        bridge._ensure_outline_claim_coverage(
            run,
            json.dumps(
                {
                    "title": "测试",
                    "slides": [
                        {
                            "layout": "title",
                            "title": "封面",
                            "source_claim_ids": ["unknown-claim"],
                        }
                    ],
                },
                ensure_ascii=False,
            ),
        )

    shorthand = json.loads(
        bridge._ensure_outline_claim_coverage(
            run,
            json.dumps(
                {
                    "title": "测试",
                    "slides": [
                        {
                            "layout": "title",
                            "title": "封面",
                            "source_claim_ids": ["C01"],
                        }
                    ],
                },
                ensure_ascii=False,
            ),
        )
    )
    assert shorthand["slides"][0]["source_claim_ids"][0] != "C01"
    assert shorthand["slides"][0]["source_claim_ids"][0].endswith(
        hashlib.sha256("第一条事实。".encode()).hexdigest()[:12]
    )


def test_bridge_outline_claim_coverage_normalizes_stale_ordinal_hash():
    import scripts.hermes_bridge as bridge
    from backend.services.presentation_source_trace import build_source_claims

    source_text = "第一条事实。第二条事实。"
    source_id = hashlib.sha256(source_text.encode("utf-8")).hexdigest()
    run = {
        "source_material": {
            "text": source_text,
            "source_id": source_id,
            "source_client_session_id": "ios-session-1",
        }
    }
    repaired = json.loads(
        bridge._ensure_outline_claim_coverage(
            run,
            json.dumps(
                {
                    "title": "测试",
                    "slides": [
                        {
                            "layout": "title",
                            "title": "封面",
                            "source_claim_ids": ["c001:9b2b75bb0f78"],
                        }
                    ],
                },
                ensure_ascii=False,
            ),
        )
    )
    claims = build_source_claims(
        source_text,
        source_id=source_id,
        source_client_session_id="ios-session-1",
        approval_state="approved",
    )
    assert repaired["slides"][0]["source_claim_ids"][0] == claims[0]["claim_id"]


def test_bridge_distributes_missing_istanbul_claims_without_appendix_slides():
    import scripts.hermes_bridge as bridge
    from backend.services.presentation_source_trace import build_source_claims

    source_text = (Path(__file__).parent / "fixtures/presentation/istanbul-source.md").read_text()
    digest = hashlib.sha256(source_text.encode()).hexdigest()
    run = {
        "source_material": {
            "text": source_text,
            "source_id": f"txt_{digest[:32]}",
            "source_client_session_id": "ios-session-1",
        }
    }
    outline = {
        "title": "伊斯坦布尔",
        "slides": [
            {"layout": "title", "title": "伊斯坦布尔两日总览", "key_points": [], "source_claim_ids": []},
            {"layout": "bullets", "title": "入境与签证", "key_points": [], "source_claim_ids": []},
            {"layout": "timeline", "title": "机场交通与换乘", "key_points": [], "source_claim_ids": []},
            {"layout": "route_map", "title": "D1 欧洲侧", "key_points": [], "source_claim_ids": []},
            {"layout": "route_map", "title": "D2 欧洲侧到亚洲侧", "key_points": [], "source_claim_ids": []},
        ],
    }
    repaired = json.loads(
        bridge._ensure_outline_claim_coverage(run, json.dumps(outline, ensure_ascii=False))
    )
    assert len(repaired["slides"]) == len(outline["slides"])
    assert all("原文完整性补充" not in slide["title"] for slide in repaired["slides"])
    claims = build_source_claims(
        source_text,
        source_id=f"txt_{digest[:32]}",
        source_client_session_id="ios-session-1",
        approval_state="approved",
    )
    bound = {
        claim_id
        for slide in repaired["slides"]
        for claim_id in slide["source_claim_ids"]
    }
    assert bound == {claim["claim_id"] for claim in claims}
    projected = "\n".join(
        point for slide in repaired["slides"] for point in slide.get("key_points") or []
    )
    assert "普通护照并非免签" in projected
    assert "不适合女生" not in projected
    assert "Kuzguncuk在亚洲侧" in projected


def test_bridge_visual_layout_normalization_is_renderer_safe_and_fail_closed():
    import scripts.hermes_bridge as bridge

    valid = {
        "slides": [{
            "layout": "icon_grid",
            "title": "Visual",
            "items": [
                {"icon": "map", "title": "Route", "detail": "Old City"},
                {"icon": "ferry", "title": "Crossing", "detail": "Bosphorus"},
            ],
        }]
    }
    normalized = json.loads(bridge._normalize_presentation_reply(json.dumps(valid)))
    assert normalized["slides"][0]["layout"] == "icon_grid"

    for mutation in ("unexpected_field", "unsupported_icon"):
        invalid = json.loads(json.dumps(valid))
        invalid["slides"][0]["bullets"] = ["Fallback"]
        if mutation == "unexpected_field":
            invalid["slides"][0]["unexpected"] = "must fail closed"
        else:
            invalid["slides"][0]["items"][0]["icon"] = "remote-svg"
        degraded = json.loads(bridge._normalize_presentation_reply(json.dumps(invalid)))
        assert degraded["slides"][0] == {
            "layout": "bullets",
            "title": "Visual",
            "bullets": ["Fallback"],
        }


def test_two_real_pptx_files_apply_distinct_theme_to_title_body_table_chart_and_section(
    tmp_path,
):
    slides = [
        {"layout": "title", "title": "主题标题", "subtitle": "副标题"},
        {"layout": "bullets", "title": "正文页", "bullets": ["正文内容"]},
        {"layout": "table", "title": "表格页", "headers": ["指标"], "rows": [["收入"]]},
        {
            "layout": "chart",
            "title": "图表页",
            "categories": ["Q1"],
            "series": [{"name": "收入", "values": [1]}],
        },
        {"layout": "section", "title": "章节页"},
    ]
    decks = []
    for name, theme in (("blue", THEME_A), ("red", THEME_B)):
        path = tmp_path / f"{name}.pptx"
        path.write_bytes(
            build_pptx(
                json.dumps(
                    {"title": name, "theme": theme, "slides": slides},
                    ensure_ascii=False,
                )
            )
        )
        decks.append(Presentation(path))
    for deck, theme in zip(decks, (THEME_A, THEME_B)):
        title_run = deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
        body_run = (
            next(
                shape
                for shape in deck.slides[1].shapes
                if getattr(shape, "has_text_frame", False) and "正文内容" in shape.text
            )
            .text_frame.paragraphs[0]
            .runs[0]
        )
        table = next(shape.table for shape in deck.slides[2].shapes if shape.has_table)
        chart = next(shape.chart for shape in deck.slides[3].shapes if shape.has_chart)
        section_run = deck.slides[4].shapes[0].text_frame.paragraphs[0].runs[0]
        assert (
            title_run.font.name == theme["fonts"]["title"]
            and str(title_run.font.color.rgb) == theme["colors"]["text"][1:]
        )
        assert (
            body_run.font.name == theme["fonts"]["body"]
            and str(body_run.font.color.rgb) == theme["colors"]["text"][1:]
        )
        assert (
            str(table.cell(0, 0).fill.fore_color.rgb) == theme["colors"]["primary"][1:]
        )
        assert (
            table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.name
            == theme["fonts"]["body"]
        )
        assert (
            str(chart.series[0].format.fill.fore_color.rgb)
            == theme["colors"]["primary"][1:]
        )
        assert (
            str(deck.slides[4].background.fill.fore_color.rgb)
            == theme["colors"]["primary"][1:]
        )
        assert (
            section_run.font.name == theme["fonts"]["title"]
            and str(section_run.font.color.rgb) == theme["colors"]["inverse"][1:]
        )
    assert (
        decks[0].slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.name
        != decks[1].slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.name
    )


def test_theme_schema_rejects_unknown_or_incomplete_fields():
    with pytest.raises(ValueError, match="unsupported"):
        build_pptx(
            json.dumps(
                {
                    "theme": {**THEME_A, "style": "ignored"},
                    "slides": [{"layout": "title", "title": "x"}],
                }
            )
        )
    incomplete = {"colors": {"primary": "#000000"}, "fonts": THEME_A["fonts"]}
    with pytest.raises(ValueError, match="incomplete"):
        build_pptx(
            json.dumps(
                {"theme": incomplete, "slides": [{"layout": "title", "title": "x"}]}
            )
        )


def test_bridge_theme_validation_stays_inside_bridge_worker_dependencies():
    from pathlib import Path

    bridge_source = Path("scripts/hermes_bridge.py").read_text()
    scenario_source = Path("backend/services/presentation_scenario.py").read_text()
    assert "presentation_scenario import validate_theme" in bridge_source
    assert "from pptx" not in scenario_source


@pytest.mark.parametrize(
    "slide",
    [
        {"layout": "bullets", "title": "x", "bullets": ["item"] * 13},
        {"layout": "bullets", "title": "x" * 181, "bullets": []},
        {
            "layout": "table",
            "title": "x",
            "headers": ["a", "b"],
            "rows": [["only one"]],
        },
        {"layout": "unknown", "title": "x"},
    ],
)
def test_presentation_renderer_rejects_content_it_cannot_render_without_truncation(
    slide,
):
    with pytest.raises(ValueError):
        build_pptx(json.dumps({"slides": [slide]}))


def test_presentation_prompt_rejects_oversized_private_source_instead_of_excerpting():
    import scripts.hermes_bridge as bridge

    node = {
        "id": "presentation_outline",
        "node_type": "LLM_INFERENCE",
        "name": "outline",
        "parameters": {"output_format": "presentation_outline", "max_tokens": 8000},
    }
    run = {
        "goal": "deck",
        "deliverable": "pptx",
        "plan": {"nodes": [node], "edges": []},
        "nodes": {},
        "source_document": {"filename": "private.docx", "text": "x" * 80001},
    }
    with pytest.raises(RuntimeError, match="禁止静默截断"):
        bridge._workflow_node_prompt(run, node)


def test_final_presentation_prompt_accepts_full_approved_context_above_chat_limit(monkeypatch):
    import scripts.hermes_bridge as bridge

    approved_outline = {
        "title": "Istanbul",
        "slides": [
            {"layout": "bullets", "title": f"slide-{index}", "key_points": ["证据" * 800]}
            for index in range(8)
        ],
    }
    monkeypatch.setattr(bridge, "_approved_presentation_outline", lambda run: (approved_outline, {}))
    design = {"id": "presentation_design", "name": "design", "parameters": {"output_format": "presentation_design"}}
    deck = {"id": "presentation_deck", "name": "deck", "node_type": "LLM_INFERENCE", "parameters": {"output_format": "presentation", "max_tokens": 8000}}
    run = {
        "goal": "deck",
        "deliverable": "pptx",
        "plan": {"nodes": [design, deck], "edges": [{"source": "presentation_design", "target": "presentation_deck"}]},
        "nodes": {"presentation_design": {"status": "succeeded", "output": json.dumps({"theme": THEME_A, "slides": []})}},
    }
    prompt = bridge._workflow_node_prompt(run, deck)
    assert bridge.MAX_INPUT < len(prompt) <= bridge.MAX_GENERATIVE_WORKFLOW_INPUT
    assert "已批准逐页大纲" in prompt


def test_presentation_design_accepts_outline_within_its_declared_output_limit(monkeypatch):
    import scripts.hermes_bridge as bridge

    outline_text = "o" * 7_000
    outline = {
        "id": "presentation_outline",
        "name": "outline",
        "node_type": "LLM_INFERENCE",
        "parameters": {"output_format": "presentation_outline", "max_tokens": 8000},
    }
    design = {
        "id": "presentation_design",
        "name": "design",
        "node_type": "LLM_INFERENCE",
        "parameters": {"output_format": "presentation_design", "max_tokens": 8000},
    }
    monkeypatch.setattr(bridge, "_approved_presentation_outline", lambda run: ({}, {}))
    run = {
        "goal": "deck",
        "deliverable": "pptx",
        "plan": {
            "nodes": [outline, design],
            "edges": [{"source": "presentation_outline", "target": "presentation_design"}],
        },
        "nodes": {
            "presentation_outline": {"status": "succeeded", "output": outline_text}
        },
    }
    prompt = bridge._workflow_node_prompt(run, design)
    assert outline_text in prompt


def test_presentation_design_accepts_full_fourteen_slide_outline(monkeypatch):
    import scripts.hermes_bridge as bridge

    outline_text = "o" * 18_000
    outline = {
        "id": "presentation_outline",
        "name": "outline",
        "node_type": "LLM_INFERENCE",
        "parameters": {"output_format": "presentation_outline", "max_tokens": 8000},
    }
    design = {
        "id": "presentation_design",
        "name": "design",
        "node_type": "LLM_INFERENCE",
        "parameters": {"output_format": "presentation_design", "max_tokens": 8000},
    }
    monkeypatch.setattr(bridge, "_approved_presentation_outline", lambda run: ({}, {}))
    run = {
        "goal": "deck",
        "deliverable": "pptx",
        "plan": {
            "nodes": [outline, design],
            "edges": [{"source": "presentation_outline", "target": "presentation_design"}],
        },
        "nodes": {"presentation_outline": {"status": "succeeded", "output": outline_text}},
    }
    prompt = bridge._workflow_node_prompt(run, design)
    assert outline_text in prompt


def test_final_presentation_uses_approved_outline_projection_without_raw_duplication(monkeypatch):
    import scripts.hermes_bridge as bridge

    approved_outline = {
        "title": "Istanbul",
        "slides": [{"layout": "route_map", "title": "D1", "key_points": ["route"]}],
    }
    monkeypatch.setattr(bridge, "_approved_presentation_outline", lambda run: (approved_outline, {}))
    outline = {"id": "presentation_outline", "name": "outline", "parameters": {"output_format": "presentation_outline"}}
    design = {"id": "presentation_design", "name": "design", "parameters": {"output_format": "presentation_design"}}
    deck = {"id": "presentation_deck", "name": "deck", "node_type": "OUTPUT_FORMAT", "parameters": {"output_format": "presentation", "max_tokens": 16000}}
    run = {
        "goal": "deck",
        "deliverable": "pptx",
        "plan": {
            "nodes": [outline, design, deck],
            "edges": [
                {"source": "presentation_outline", "target": "presentation_deck"},
                {"source": "presentation_design", "target": "presentation_deck"},
            ],
        },
        "nodes": {
            "presentation_outline": {"status": "succeeded", "output": "RAW_OUTLINE_MUST_NOT_REPEAT"},
            "presentation_design": {"status": "succeeded", "output": "approved design"},
        },
    }
    prompt = bridge._workflow_node_prompt(run, deck)
    assert "RAW_OUTLINE_MUST_NOT_REPEAT" not in prompt
    assert "已批准逐页大纲" in prompt
    assert "approved design" in prompt


def test_final_presentation_prompt_still_rejects_context_above_generation_limit(monkeypatch):
    import scripts.hermes_bridge as bridge

    monkeypatch.setattr(
        bridge,
        "_approved_presentation_outline",
        lambda run: ({"title": "oversized", "slides": [{"layout": "bullets", "title": "x", "key_points": ["证据" * 17_000]}]}, {}),
    )
    deck = {"id": "presentation_deck", "name": "deck", "node_type": "LLM_INFERENCE", "parameters": {"output_format": "presentation", "max_tokens": 8000}}
    run = {"goal": "deck", "deliverable": "pptx", "plan": {"nodes": [deck], "edges": []}, "nodes": {}}
    with pytest.raises(RuntimeError, match="超过 32000 字符上限"):
        bridge._workflow_node_prompt(run, deck)


def test_bridge_normalizes_outline_style_points_for_renderable_slides():
    import scripts.hermes_bridge as bridge

    normalized = json.loads(bridge._normalize_presentation_reply(json.dumps({
        "title": "design",
        "theme": THEME_A,
        "slides": [
            {"layout": "title", "title": "cover", "subtitle": "brief", "bullets": ["unused"], "purpose": "cover"},
            {"layout": "bullets", "title": "goals", "key_points": ["one"]},
            {"layout": "section", "title": "flow", "subtitle": "one → two"},
            {"layout": "section", "title": "checkpoints", "key_points": ["upload", "preview"]},
            {"layout": "process", "title": "steps", "steps": ["one", "two"]},
            {"layout": "two_column", "title": "governance", "left_title": "data", "left_points": ["private"], "right_title": "flow", "right_points": ["approve"]},
            {"layout": "two_column", "title": "nested", "left": {"heading": "data", "points": ["private"]}, "right": {"heading": "flow", "points": ["approve"]}},
            {"layout": "two_column", "title": "empty columns", "subtitle": "preserve subtitle"},
            {"layout": "chart", "title": "no verified data"},
            {"layout": "table", "title": "fallback points", "key_points": ["one", "two"]},
        ],
    })))
    assert normalized["slides"][0] == {"layout": "title", "title": "cover", "subtitle": "brief"}
    assert normalized["slides"][1]["bullets"] == ["one"]
    assert "key_points" not in normalized["slides"][1]
    assert normalized["slides"][2]["layout"] == "title"
    assert normalized["slides"][3]["layout"] == "bullets"
    assert normalized["slides"][3]["bullets"] == ["upload", "preview"]
    assert normalized["slides"][4]["layout"] == "bullets"
    assert normalized["slides"][4]["bullets"] == ["one", "two"]
    assert normalized["slides"][5]["left"] == ["data", "private"]
    assert normalized["slides"][5]["right"] == ["flow", "approve"]
    assert "left_points" not in normalized["slides"][5]
    assert "right_title" not in normalized["slides"][5]
    assert normalized["slides"][6]["left"] == ["data", "private"]
    assert normalized["slides"][6]["right"] == ["flow", "approve"]
    assert "steps" not in normalized["slides"][4]
    assert normalized["slides"][7] == {"layout": "title", "title": "empty columns", "subtitle": "preserve subtitle"}
    assert normalized["slides"][8] == {"layout": "section", "title": "no verified data"}
    assert normalized["slides"][9] == {"layout": "bullets", "title": "fallback points", "bullets": ["one", "two"]}
    build_pptx(json.dumps(normalized))


def test_bridge_final_projection_uses_exact_approved_design_theme_and_binding():
    import scripts.hermes_bridge as bridge

    design = json.dumps(
        {
            "title": "approved",
            "theme": THEME_A,
            "slides": [{"layout": "title", "title": "sample"}],
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    digest = hashlib.sha256(design.encode()).hexdigest()
    run = {
        "plan": {
            "nodes": [
                {
                    "id": "presentation_design",
                    "parameters": {"output_format": "presentation_design"},
                }
            ]
        },
        "nodes": {
            "presentation_design": {
                "status": "succeeded",
                "attempt": 3,
                "output": design,
            }
        },
        "approved_gates": ["presentation_design"],
        "approved_gate_artifacts": {
            "presentation_design": {
                "artifact_id": "wfa_" + "a" * 32,
                "content_hash": digest,
                "artifact_version": 3,
            }
        },
    }
    model_final = json.dumps(
        {
            "theme": THEME_B,
            "slides": [{"layout": "bullets", "title": "final", "bullets": ["kept"]}],
        }
    )
    projected, binding = bridge._bind_approved_presentation_design(run, model_final)
    assert json.loads(projected)["theme"] == {
        "colors": {key: value.upper() for key, value in THEME_A["colors"].items()},
        "fonts": THEME_A["fonts"],
    }
    assert (
        json.loads(projected)["slides"][0]["bullets"] == ["kept"]
        and binding["content_hash"] == digest
    )
    run["nodes"]["presentation_design"]["output"] = design + " "
    with pytest.raises(RuntimeError, match="tampered"):
        bridge._bind_approved_presentation_design(run, model_final)


@pytest.mark.asyncio
async def test_workflow_projection_rechecks_approved_design_against_final_theme(
    tmp_path, monkeypatch
):
    from types import SimpleNamespace
    from backend.services.workflow_executor import _assert_approved_presentation_projection
    import backend.services.workflow_artifacts as artifacts

    approved = json.dumps(
        {"theme": THEME_A, "slides": [{"layout": "title", "title": "sample"}]}
    )
    path = tmp_path / "design.json"
    path.write_text(approved)
    digest = hashlib.sha256(approved.encode()).hexdigest()
    design = SimpleNamespace(
        id="wfa_" + "a" * 32,
        execution_id="exec-theme",
        content_hash=digest,
        relative_path="design.json",
    )
    outline_value = {
        "title": "final",
        "slides": [{"layout": "title", "title": "final"}],
    }
    outline_text = json.dumps(outline_value)
    outline_path = tmp_path / "outline.json"
    outline_path.write_text(outline_text)
    outline_digest = hashlib.sha256(outline_text.encode()).hexdigest()
    outline = SimpleNamespace(
        id="wfa_" + "b" * 32,
        execution_id="exec-theme",
        content_hash=outline_digest,
        relative_path="outline.json",
    )

    class DB:
        async def scalar(self, query):
            return object()

        async def get(self, model, key):
            return {design.id: design, outline.id: outline}.get(key)

    monkeypatch.setattr(artifacts, "run_root", lambda execution: tmp_path)
    binding = {"artifact_id": design.id, "content_hash": digest, "artifact_version": 2}
    outline_binding = {
        "artifact_id": outline.id,
        "content_hash": outline_digest,
        "artifact_version": 1,
    }
    execution = SimpleNamespace(id="exec-theme")
    await _assert_approved_presentation_projection(
        DB(),
        execution,
        {
            "approved_design": binding,
            "approved_outline": outline_binding,
            "content": json.dumps(
                {"theme": THEME_A, "slides": [{"layout": "title", "title": "final"}]}
            ),
        },
        "presentation",
    )
    await _assert_approved_presentation_projection(
        DB(),
        execution,
        {
            "approved_design": binding,
            "approved_outline": outline_binding,
            "content": json.dumps(
                {"theme": THEME_A, "slides": [{"layout": "bullets", "title": "final", "bullets": ["safe fallback"]}]}
            ),
        },
        "presentation",
    )
    with pytest.raises(ValueError, match="differs"):
        await _assert_approved_presentation_projection(
            DB(),
            execution,
            {
                "approved_design": binding,
                "approved_outline": outline_binding,
                "content": json.dumps(
                    {
                        "theme": THEME_B,
                        "slides": [{"layout": "title", "title": "tampered"}],
                    }
                ),
            },
            "presentation",
        )


@pytest.mark.asyncio
async def test_workflow_projection_accepts_only_matching_ungated_stage_artifacts(
    tmp_path, monkeypatch
):
    from types import SimpleNamespace
    import backend.services.workflow_artifacts as artifacts
    import backend.services.workflow_executor as executor

    design_text = json.dumps({
        "theme": THEME_A,
        "slides": [{"layout": "title", "title": "sample"}],
    })
    outline_text = json.dumps({
        "title": "deck",
        "slides": [{"layout": "title", "title": "deck"}],
    })
    (tmp_path / "design.json").write_text(design_text)
    (tmp_path / "outline.json").write_text(outline_text)
    design_hash = hashlib.sha256(design_text.encode()).hexdigest()
    outline_hash = hashlib.sha256(outline_text.encode()).hexdigest()
    design_node = SimpleNamespace(id="node-design", node_id="presentation_design")
    outline_node = SimpleNamespace(id="node-outline", node_id="presentation_outline")
    design = SimpleNamespace(
        execution_id="exec-generated", node_run_id=design_node.id,
        content_hash=design_hash, relative_path="design.json",
        metadata_json={"artifact_version": 1, "render_type": "presentation_design"},
    )
    outline = SimpleNamespace(
        execution_id="exec-generated", node_run_id=outline_node.id,
        content_hash=outline_hash, relative_path="outline.json",
        metadata_json={"artifact_version": 1, "render_type": "presentation_outline"},
    )

    class DB:
        def __init__(self):
            self.artifacts = iter([design, outline])

        async def scalar(self, query):
            return next(self.artifacts)

    monkeypatch.setattr(artifacts, "run_root", lambda execution: tmp_path)
    monkeypatch.setattr(
        executor,
        "_plan",
        lambda db, execution: asyncio.sleep(0, result=SimpleNamespace(dsl={"nodes": [
            {"id": "presentation_design", "parameters": {"output_format": "presentation_design"}},
            {"id": "presentation_outline", "parameters": {"output_format": "presentation_outline"}},
        ]})),
    )
    monkeypatch.setattr(
        executor,
        "_nodes",
        lambda db, execution_id: asyncio.sleep(0, result={
            design_node.node_id: design_node,
            outline_node.node_id: outline_node,
        }),
    )
    execution = SimpleNamespace(id="exec-generated", plan_id="plan-generated")
    final = {
        "theme": THEME_A,
        "slides": [{"layout": "title", "title": "deck"}],
    }
    await executor._assert_approved_presentation_projection(
        DB(), execution,
        {
            "approved_design": {
                "approval_mode": "generated_verified", "node_id": design_node.node_id,
                "content_hash": design_hash, "artifact_version": 1,
            },
            "approved_outline": {
                "approval_mode": "generated_verified", "node_id": outline_node.node_id,
                "content_hash": outline_hash, "artifact_version": 1,
            },
            "content": json.dumps(final),
        },
        "presentation",
    )


def test_final_presentation_requires_approved_outline_binding_and_structure():
    import scripts.hermes_bridge as bridge

    design = json.dumps({"theme": THEME_A, "slides": [{"layout": "title", "title": "sample"}]})
    outline = json.dumps(
        {
            "title": "deck",
            "slides": [
                {"layout": "title", "title": "deck"},
                {"layout": "bullets", "title": "findings", "key_points": ["one"]},
            ],
        }
    )
    run = {
        "plan": {
            "nodes": [
                {"id": "presentation_outline", "parameters": {"output_format": "presentation_outline"}},
                {"id": "presentation_design", "parameters": {"output_format": "presentation_design"}},
            ]
        },
        "nodes": {
            "presentation_outline": {"attempt": 1, "output": outline},
            "presentation_design": {"attempt": 2, "output": design},
        },
        "approved_gates": ["presentation_outline", "presentation_design"],
        "approved_gate_artifacts": {
            "presentation_outline": {
                "artifact_id": "wfa_" + "a" * 32,
                "content_hash": hashlib.sha256(outline.encode()).hexdigest(),
                "artifact_version": 1,
            },
            "presentation_design": {
                "artifact_id": "wfa_" + "b" * 32,
                "content_hash": hashlib.sha256(design.encode()).hexdigest(),
                "artifact_version": 2,
            },
        },
    }
    final = json.dumps(
        {
            "title": "deck",
            "slides": [
                {"layout": "title", "title": "deck"},
                {"layout": "bullets", "title": "findings", "bullets": ["one"]},
            ],
        }
    )
    projected, design_binding, outline_binding = bridge._bind_approved_presentation_inputs(run, final)
    assert json.loads(projected)["theme"]["fonts"] == THEME_A["fonts"]
    assert design_binding["artifact_version"] == 2
    assert outline_binding["artifact_version"] == 1
    changed = json.loads(final)
    changed["slides"][1]["title"] = "unapproved"
    rebound, _, _ = bridge._bind_approved_presentation_inputs(run, json.dumps(changed))
    assert json.loads(rebound)["slides"][1]["title"] == json.loads(outline)["slides"][1]["title"]
    shortened = json.loads(final)
    shortened["slides"] = shortened["slides"][:1]
    rebuilt, _, _ = bridge._bind_approved_presentation_inputs(run, json.dumps(shortened))
    rebuilt_slides = json.loads(rebuilt)["slides"]
    assert len(rebuilt_slides) == len(json.loads(outline)["slides"])
    assert rebuilt_slides[1]["layout"] == "bullets"
    assert rebuilt_slides[1]["bullets"] == json.loads(outline)["slides"][1]["key_points"]
    run["nodes"]["presentation_outline"]["output"] = outline + " "
    with pytest.raises(RuntimeError, match="outline.*tampered"):
        bridge._bind_approved_presentation_inputs(run, final)


def test_default_presentation_binds_verified_ungated_outline_and_design():
    import scripts.hermes_bridge as bridge

    outline = json.dumps({
        "title": "deck",
        "slides": [{"layout": "title", "title": "deck"}],
    })
    design = json.dumps({
        "theme": THEME_A,
        "slides": [{"layout": "title", "title": "sample"}],
    })
    run = {
        "plan": {"nodes": [
            {"id": "presentation_outline", "parameters": {"output_format": "presentation_outline"}},
            {"id": "presentation_design", "parameters": {"output_format": "presentation_design"}},
        ]},
        "nodes": {
            "presentation_outline": {
                "status": "succeeded", "attempt": 1, "output": outline,
                "content_hash": hashlib.sha256(outline.encode()).hexdigest(),
            },
            "presentation_design": {
                "status": "succeeded", "attempt": 1, "output": design,
                "content_hash": hashlib.sha256(design.encode()).hexdigest(),
            },
        },
        "approved_gates": [],
        "approved_gate_artifacts": {},
    }
    final = json.dumps({
        "title": "deck",
        "slides": [{"layout": "title", "title": "deck"}],
    })

    _, design_binding, outline_binding = bridge._bind_approved_presentation_inputs(run, final)

    assert design_binding == {
        "approval_mode": "generated_verified",
        "node_id": "presentation_design",
        "content_hash": hashlib.sha256(design.encode()).hexdigest(),
        "artifact_version": 1,
    }
    assert outline_binding == {
        "approval_mode": "generated_verified",
        "node_id": "presentation_outline",
        "content_hash": hashlib.sha256(outline.encode()).hexdigest(),
        "artifact_version": 1,
    }
    run["nodes"]["presentation_outline"]["output"] = outline + " "
    with pytest.raises(RuntimeError, match="outline.*tampered"):
        bridge._bind_approved_presentation_inputs(run, final)


def test_hermes_gate_rejects_stale_version_and_records_current_approval(monkeypatch):
    import scripts.hermes_bridge as bridge
    from fastapi import HTTPException

    output = '{"title":"outline","slides":[{"layout":"title","title":"x"}]}'
    digest = hashlib.sha256(output.encode()).hexdigest()
    run = {
        "execution_id": "exec-gate",
        "status": "awaiting_approval",
        "next_seq": 1,
        "events": [],
        "nodes": {"outline": {"status": "succeeded", "attempt": 2, "output": output}},
        "approved_gates": [],
        "approved_gate_artifacts": {},
        "plan": {
            "nodes": [{"id": "outline", "parameters": {"approval_gate": "outline"}}],
            "edges": [],
        },
    }
    monkeypatch.setattr(bridge, "HERMES_BRIDGE_INTERNAL_TOKEN", "secret")
    monkeypatch.setattr(bridge, "_start_workflow_thread", lambda execution_id: None)
    monkeypatch.setattr(bridge, "_save_workflow_runs", lambda: None)
    bridge._workflow_runs["exec-gate"] = run
    try:
        with pytest.raises(HTTPException) as stale:
            asyncio.run(
                bridge.approve_workflow_gate(
                    "exec-gate",
                    bridge.WorkflowGateApprovalRequest(
                        node_id="outline",
                        artifact_version=1,
                        artifact_id="wfa_" + "a" * 32,
                        expected_hash=digest,
                    ),
                    "secret",
                )
            )
        assert stale.value.status_code == 409
        with pytest.raises(HTTPException) as tampered:
            asyncio.run(
                bridge.approve_workflow_gate(
                    "exec-gate",
                    bridge.WorkflowGateApprovalRequest(
                        node_id="outline",
                        artifact_version=2,
                        artifact_id="wfa_" + "a" * 32,
                        expected_hash="0" * 64,
                    ),
                    "secret",
                )
            )
        assert tampered.value.status_code == 409
        result = asyncio.run(
            bridge.approve_workflow_gate(
                "exec-gate",
                bridge.WorkflowGateApprovalRequest(
                    node_id="outline",
                    artifact_version=2,
                    artifact_id="wfa_" + "a" * 32,
                    expected_hash=digest,
                ),
                "secret",
            )
        )
        assert result["status"] == "queued" and run["approved_gates"] == ["outline"]
        assert run["approved_gate_artifacts"]["outline"]["content_hash"] == digest
        retried = asyncio.run(
            bridge.retry_workflow_run(
                "exec-gate",
                bridge.WorkflowRetryRequest(
                    from_node_id="outline",
                    revision_comment="第二页标题需要更具体",
                ),
                "secret",
            )
        )
        assert (
            retried["status"] == "queued"
            and run["approved_gates"] == []
            and run["approved_gate_artifacts"] == {}
        )
        assert "第二页标题需要更具体" in bridge._workflow_node_prompt(
            run, run["plan"]["nodes"][0]
        )
    finally:
        bridge._workflow_runs.pop("exec-gate", None)


def test_workflow_retry_refreshes_expired_authorization_without_scope_widening(monkeypatch):
    import scripts.hermes_bridge as bridge

    run = {
        "execution_id": "exec-refresh",
        "tenant_id": "tenant-a",
        "knowledge_capability": "expired-capability",
        "knowledge_policy_version": "old-policy",
        "knowledge_scope": ["private-old"],
    }
    monkeypatch.setattr(
        bridge,
        "_validated_knowledge_claims",
        lambda token, **_: {
            "entry_point": "workflow",
            "tenant_key": "tenant-a",
            "scopes": [],
        },
    )
    bridge._refresh_workflow_authorization(
        run, "replacement-capability-value", "replacement-policy"
    )
    assert run["knowledge_capability"] == "replacement-capability-value"
    assert run["knowledge_policy_version"] == "replacement-policy"
    assert run["knowledge_scope"] == []


def test_presentation_scenario_defaults_to_full_draft_review_and_binary_deck():
    workflow = type(
        "Workflow",
        (),
        {
            "title": "Deck",
            "requirements_snapshot": {
                "scenario_id": "presentation-generation",
                "source_document": {
                    "source_id": "doc_1",
                    "content_hash": "a" * 64,
                    "source_revision": 1,
                },
            },
        },
    )()
    plan = build_presentation_plan(workflow, plan_id="plan", knowledge_scope=[])
    assert plan is not None
    assert [node["parameters"].get("approval_gate") for node in plan["nodes"]] == [
        None,
        None,
        None,
        None,
    ]
    assert plan["version"] == "3.0.0"
    assert plan["nodes"][0]["id"] == "presentation_analysis"
    assert plan["nodes"][-1]["parameters"]["output_format"] == "presentation"
    assert {tuple(edge.values()) for edge in plan["edges"]} >= {
        ("presentation_analysis", "presentation_outline"),
        ("presentation_outline", "presentation_deck"),
        ("presentation_design", "presentation_deck"),
    }
    assert len(DSLSafetyCompiler.compile_and_validate(plan).nodes) == 4


def test_presentation_risk_review_gates_require_explicit_policy():
    snapshot = {
        "scenario_id": "presentation-generation",
        "text_material": "Approved source material.",
        "presentation_review_gates": ["outline", "design"],
    }
    workflow = type(
        "Workflow",
        (),
        {
            "title": "Regulated deck",
            "requirements_snapshot": snapshot,
        },
    )()
    plan = build_presentation_plan(workflow, plan_id="plan", knowledge_scope=[])
    assert plan is not None
    assert [node["parameters"].get("approval_gate") for node in plan["nodes"]] == [
        None,
        "outline",
        "design",
        None,
    ]

    snapshot["presentation_review_gates"] = ["unknown"]
    with pytest.raises(ValueError, match="presentation_review_gates"):
        build_presentation_plan(workflow, plan_id="plan", knowledge_scope=[])


def test_presentation_default_path_has_no_intermediate_approval_gates():
    workflow = type(
        "Workflow",
        (),
        {
            "title": "Default three-step deck",
            "description": "Requirements are already confirmed.",
            "requirements_snapshot": {
                "scenario_id": "presentation-generation",
                "text_material": "Approved source material.",
            },
        },
    )()

    plan = build_presentation_plan(workflow, plan_id="plan", knowledge_scope=[])

    assert plan is not None
    assert [
        node["parameters"].get("approval_gate") for node in plan["nodes"]
    ] == [None, None, None, None]


def test_presentation_scenario_without_upload_researches_user_topic_first():
    workflow = type(
        "Workflow",
        (),
        {
            "title": "鹿儿岛旅行攻略",
            "description": "帮我生成一个介绍鹿儿岛旅行攻略的 PPT",
            "requirements_snapshot": {"scenario_id": "presentation-generation"},
        },
    )()
    plan = build_presentation_plan(workflow, plan_id="plan", knowledge_scope=[])
    assert plan["source_document"] == {}
    assert plan["nodes"][0]["id"] == "presentation_research"
    assert plan["nodes"][0]["node_type"] == "KNOWLEDGE_RETRIEVAL"
    assert plan["nodes"][0]["parameters"]["knowledge_scope"] == []
    assert {tuple(edge.values()) for edge in plan["edges"]} >= {
        ("presentation_research", "presentation_analysis"),
    }


def test_text_material_presentation_does_not_read_unscoped_private_knowledge():
    workflow = type(
        "Workflow",
        (),
        {
            "title": "因特拉肯管理层简报",
            "description": "根据用户材料生成六页管理层 PPT",
            "requirements_snapshot": {
                "scenario_id": "presentation-generation",
                "text_material": "因特拉肯位于图恩湖与布里恩茨湖之间。",
            },
        },
    )()
    plan = build_presentation_plan(workflow, plan_id="plan", knowledge_scope=[])
    assert plan is not None
    assert plan["source_document"] == {}
    assert "presentation_research" not in {node["id"] for node in plan["nodes"]}
    assert ("presentation_research", "presentation_analysis") not in {
        tuple(edge.values()) for edge in plan["edges"]
    }
    assert "用户提供的文字材料" in plan["nodes"][0]["parameters"]["instruction"]


def test_inline_text_material_is_hash_bound_to_its_client_session():
    from backend.services.workflow_executor import inline_source_material

    material = inline_source_material(
        {
            "text_material": "只允许当前会话使用的伊斯坦布尔材料。",
            "source_client_session_id": "session-istanbul",
        },
        workflow_id="wf-current",
    )
    assert material is not None
    assert material["source_id"].startswith("txt_")
    assert material["source_client_session_id"] == "session-istanbul"
    assert material["scope_state"] == "bound"
    assert material["content_hash"] == hashlib.sha256(
        material["text"].encode("utf-8")
    ).hexdigest()

    legacy = inline_source_material(
        {"text_material": "升级前材料"}, workflow_id="wf-legacy"
    )
    assert legacy is not None
    assert legacy["source_client_session_id"] == "legacy-workflow:wf-legacy"
    assert legacy["scope_state"] == "legacy_isolated"


def test_bridge_injects_exact_claim_inventory_for_inline_presentation_material():
    import scripts.hermes_bridge as bridge

    text = "伊斯坦布尔横跨欧亚。博斯普鲁斯海峡连接两岸。"
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    analysis = {
        "id": "presentation_analysis",
        "node_type": "LLM_INFERENCE",
        "name": "分析",
        "parameters": {"output_format": "markdown", "max_tokens": 5000},
    }
    deck = {
        "id": "presentation_deck",
        "node_type": "OUTPUT_FORMAT",
        "name": "成品",
        "parameters": {"output_format": "presentation", "max_tokens": 16000},
    }
    payload = {
        "tenant_id": "tenant-a",
        "execution_id": "exec-a",
        "idempotency_key": "request-a",
        "goal": "制作伊斯坦布尔 PPT",
        "deliverable": "pptx",
        "plan": {"nodes": [analysis, deck], "edges": []},
        "knowledge_capability": "capability-token-value",
        "knowledge_policy_version": "policy-v1",
        "source_material": {
            "source_id": f"txt_{digest[:32]}",
            "content_hash": digest,
            "text": text,
            "source_client_session_id": "session-istanbul",
            "scope_state": "bound",
        },
    }
    run = bridge.WorkflowRunRequest.model_validate(payload).model_dump(exclude_none=True)
    prompt = bridge._workflow_node_prompt(run, analysis)
    assert text in prompt
    assert "批准的用户材料事实清单" in prompt
    assert f"txt_{digest[:32]}:c001" in prompt
    assert "不得引入未批准事实" in prompt

    tampered = copy.deepcopy(payload)
    tampered["source_material"]["content_hash"] = "0" * 64
    with pytest.raises(ValueError, match="invalid inline source material"):
        bridge.WorkflowRunRequest.model_validate(tampered)


def test_bridge_injects_deterministic_istanbul_travel_contract():
    import scripts.hermes_bridge as bridge

    text = (Path(__file__).parent / "fixtures/presentation/istanbul-source.md").read_text()
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    analysis = {
        "id": "presentation_analysis",
        "node_type": "LLM_INFERENCE",
        "name": "分析",
        "parameters": {"output_format": "markdown", "max_tokens": 5000},
    }
    deck = {
        "id": "presentation_deck",
        "node_type": "OUTPUT_FORMAT",
        "name": "成品",
        "parameters": {"output_format": "presentation", "max_tokens": 16000},
    }
    run = bridge.WorkflowRunRequest.model_validate(
        {
            "tenant_id": "tenant-a",
            "execution_id": "exec-istanbul-contract",
            "idempotency_key": "request-istanbul-contract",
            "goal": "制作伊斯坦布尔 PPT",
            "deliverable": "pptx",
            "plan": {"nodes": [analysis, deck], "edges": []},
            "knowledge_capability": "capability-token-value",
            "knowledge_policy_version": "policy-v1",
            "source_material": {
                "source_id": f"txt_{digest[:32]}",
                "content_hash": digest,
                "text": text,
                "source_client_session_id": "session-istanbul",
                "scope_state": "bound",
            },
        }
    ).model_dump(exclude_none=True)

    prompt = bridge._workflow_node_prompt(run, analysis)
    assert "确定性旅行编排合同" in prompt
    assert "禁止新增原文没有的酒店榜单" in prompt
    assert "中国大陆普通护照旅游或商务不是免签" in prompt
    assert "M11到Gayrettepe，再换M2" in prompt
    assert "Kuzguncuk在亚洲侧" in prompt
    assert "165里拉制卡费" in prompt


def test_bridge_corrects_m11_first_stop_wording_in_final_deck():
    import scripts.hermes_bridge as bridge

    run = {
        "source_material": {
            "text": "伊斯坦布尔 IST 搭乘M11第一站坐到Gayrettepe",
        }
    }
    value = {
        "slides": [
            {
                "title": "机场进城",
                "body": "M11第一站下车，换乘M2",
                "bullets": ["搭乘M11第一站坐到Gayrettepe"],
            }
        ]
    }

    bridge._enforce_istanbul_final_facts(run, value)

    rendered = json.dumps(value, ensure_ascii=False)
    assert "M11第一站" not in rendered
    assert "M11至Gayrettepe（并非IST后的下一站）" in rendered


def test_final_presentation_source_trace_requires_complete_approved_outline_mapping():
    import scripts.hermes_bridge as bridge
    from backend.services.presentation_source_trace import build_source_claims

    text = "伊斯坦布尔横跨欧亚。博斯普鲁斯海峡连接两岸。"
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    claims = build_source_claims(
        text,
        source_id=f"txt_{digest[:32]}",
        source_client_session_id="session-istanbul",
        approval_state="approved",
    )
    outline_value = {
        "title": "伊斯坦布尔",
        "slides": [
            {
                "layout": "title",
                "title": "横跨欧亚",
                "source_claim_ids": [claim["claim_id"] for claim in claims],
            }
        ],
    }
    outline = json.dumps(outline_value, ensure_ascii=False, separators=(",", ":"))
    run = {
        "source_material": {
            "source_id": f"txt_{digest[:32]}",
            "content_hash": digest,
            "text": text,
            "source_client_session_id": "session-istanbul",
            "scope_state": "bound",
        },
        "plan": {
            "nodes": [
                {
                    "id": "presentation_outline",
                    "parameters": {"output_format": "presentation_outline"},
                }
            ]
        },
        "nodes": {
            "presentation_outline": {
                "status": "succeeded",
                "attempt": 1,
                "output": outline,
            }
        },
        "approved_gates": ["presentation_outline"],
        "approved_gate_artifacts": {
            "presentation_outline": {
                "artifact_version": 1,
                "content_hash": hashlib.sha256(outline.encode()).hexdigest(),
            }
        },
    }
    manifest = bridge._presentation_source_trace(run)
    assert manifest is not None
    from jsonschema import Draft202012Validator

    schema = json.loads(
        Path("backend/contracts/presentation/source-trace.schema.json").read_text(
            encoding="utf-8"
        )
    )
    Draft202012Validator(schema).validate(manifest)
    assert manifest["record_count"] == len(claims)
    assert manifest["slide_ids"] == ["slide-001"]
    assert all(item["approval_state"] == "approved" for item in manifest["records"])

    incomplete = copy.deepcopy(run)
    incomplete_outline = copy.deepcopy(outline_value)
    incomplete_outline["slides"][0]["source_claim_ids"] = [claims[0]["claim_id"]]
    incomplete_text = json.dumps(incomplete_outline, ensure_ascii=False, separators=(",", ":"))
    incomplete["nodes"]["presentation_outline"]["output"] = incomplete_text
    incomplete["approved_gate_artifacts"]["presentation_outline"]["content_hash"] = (
        hashlib.sha256(incomplete_text.encode()).hexdigest()
    )
    with pytest.raises(RuntimeError, match="does not cover every source claim"):
        bridge._presentation_source_trace(incomplete)

    unmapped = copy.deepcopy(run)
    unmapped_outline = copy.deepcopy(outline_value)
    unmapped_outline["slides"].append(
        {"layout": "content", "title": "无来源页面", "source_claim_ids": []}
    )
    unmapped_text = json.dumps(unmapped_outline, ensure_ascii=False, separators=(",", ":"))
    unmapped["nodes"]["presentation_outline"]["output"] = unmapped_text
    unmapped["approved_gate_artifacts"]["presentation_outline"]["content_hash"] = (
        hashlib.sha256(unmapped_text.encode()).hexdigest()
    )
    with pytest.raises(RuntimeError, match="slides without source claims"):
        bridge._presentation_source_trace(unmapped)


def test_artifact_storage_contract_preserves_source_trace_metadata():
    from backend.services.workflow_executor import artifact_storage_contract

    node = type(
        "Node", (),
        {"node_id": "presentation_deck", "agent_id": "main", "model_used": "m", "provider_used": "p", "attempt": 1},
    )()
    trace = {"schema_version": 1, "content_hash": "a" * 64, "records": [{"claim_id": "c1"}]}
    extension, metadata = artifact_storage_contract(
        {"render_type": "presentation", "source_trace": trace},
        event_id="event-1",
        node=node,  # type: ignore[arg-type]
    )
    assert extension == "pptx"
    assert metadata["source_trace"] == trace


def test_document_scenario_generates_real_word_output_with_two_confirmation_gates():
    from backend.services.presentation_scenario import build_document_plan

    workflow = type(
        "Workflow",
        (),
        {
            "title": "项目说明书",
            "description": "帮我写一份项目说明 Word 文档",
            "requirements_snapshot": {"scenario_id": "document-generation"},
        },
    )()
    plan = build_document_plan(workflow, plan_id="plan", knowledge_scope=[])
    assert [node["parameters"].get("approval_gate") for node in plan["nodes"]] == [
        None,
        None,
        "outline",
        "content",
        None,
    ]
    assert plan["nodes"][-1]["parameters"]["output_format"] == "word"
    assert len(DSLSafetyCompiler.compile_and_validate(plan).nodes) == 5


def test_document_scenario_uses_bound_text_without_unrelated_knowledge_retrieval():
    from backend.services.presentation_scenario import build_document_plan

    workflow = type(
        "Workflow",
        (),
        {
            "title": "经营计划",
            "description": "根据已提供材料生成 Word",
            "requirements_snapshot": {
                "scenario_id": "document-generation",
                "text_material": "收入目标与交付里程碑均已由用户提供。",
                "document_profile": {"kind": "word", "evidence_policy": "user_material_only"},
            },
        },
    )()

    plan = build_document_plan(workflow, plan_id="plan", knowledge_scope=[])
    assert plan is not None

    assert [node["id"] for node in plan["nodes"]] == [
        "document_analysis",
        "document_outline",
        "document_draft",
        "document_file",
    ]
    assert all(node["node_type"] != "KNOWLEDGE_RETRIEVAL" for node in plan["nodes"])
    assert len(DSLSafetyCompiler.compile_and_validate(plan).nodes) == 4


def test_presentation_workflow_uses_presentation_questions_and_reads_source_in_analysis():
    import scripts.hermes_bridge as bridge
    from backend.api.workflows import clarification_payload, requirement_confirmation_payload

    workflow = type(
        "Workflow",
        (),
        {
            "title": "Deck",
            "description": "把文档做成 PPT",
            "desired_output": "可编辑 PPTX",
                "requirements_snapshot": {
                    "scenario_id": "presentation-generation",
                    "source_document": {"source_id": "doc_1"},
                    "source_document_evidence": "收入证明；月收入 20,000 元",
                },
        },
    )()
    first_question = clarification_payload(0, workflow)
    assert first_question["dimension"] == "用途与受众"
    assert "月收入 20,000 元" in first_question["question"]
    confirmation = requirement_confirmation_payload(workflow, [])
    assert confirmation["choices"][0] == "确认，开始生成演示文稿"
    assert "文档依据：收入证明" in confirmation["question"]

    plan = build_presentation_plan(workflow, plan_id="plan", knowledge_scope=[])
    run = {
        "goal": "deck",
        "deliverable": "pptx",
        "plan": plan,
        "nodes": {},
        "source_document": {"filename": "private.docx", "text": "只应进入分析节点的私有原文"},
    }
    analysis_prompt = bridge._workflow_node_prompt(run, plan["nodes"][0])
    outline_prompt = bridge._workflow_node_prompt(run, plan["nodes"][1])
    assert "只应进入分析节点的私有原文" in analysis_prompt
    assert "只应进入分析节点的私有原文" not in outline_prompt
