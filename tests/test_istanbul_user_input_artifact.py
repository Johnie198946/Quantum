import json
import hashlib
import re
import shutil

import pytest
from pptx import Presentation
from pypdf import PdfReader

from scripts.build_istanbul_user_input_presentation import build_package


def _all_text(path):
    presentation = Presentation(path)
    return "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text_frame") and shape.has_text_frame
    )


def test_user_input_artifact_covers_maps_lodging_food_and_source_boundaries(tmp_path):
    if not (shutil.which("soffice") or shutil.which("libreoffice")):
        pytest.skip("LibreOffice unavailable")
    pytest.importorskip("pypdfium2")

    manifest = build_package(tmp_path)
    assert manifest["visual_gate_passed"] is True
    assert manifest["input_sha256"] == "7f817d92d614362b12b928d8547f90018397793acb0c8484ba11610111ca3440"
    assert manifest["input_raw_sha256"] == "fc210e90867e714d1264ca75409cc31451d30c3025aa5b53f235cce7bd27359d"
    assert manifest["input_hash_canonicalization"] == "utf-8 text with trailing newlines removed"
    assert all(manifest["coverage"].values())

    text = _all_text(tmp_path / "istanbul-user-input.pptx")
    for required in (
        "M11 → M2 → T1",
        "Passport Control",
        "Yandex Metro",
        "D1 欧洲区 · 原文 9 站地图",
        "D2 纠偏地图",
        "Seven Hills",
        "Sirkeci Lokantası 1912",
        "Nusr-Et Kapalıçarşı",
        "Galata Konak Cafe",
        "Kuzguncuk",
        "原文没有更多酒店实测",
    ):
        assert required in text
    for generic_rewrite in ("Turkey’s largest city", "Three names, two empires"):
        assert generic_rewrite not in text

    trace = json.loads((tmp_path / "source-trace.json").read_text())
    assert trace["record_count"] >= 42
    assert trace["slide_ids"] == [f"slide-{index:03d}" for index in range(1, 13)]
    source_claims = "\n".join(record["claim"] for record in trace["records"])
    assert "黄昏轮渡" in source_claims
    assert "Kadırga Hamamı/Cemberlitas Hammam" in source_claims

    external = json.loads((tmp_path / "external-source-manifest.json").read_text())
    assert external["sources"][0]["kind"] == "official-primary"
    assert "evisa" in external["sources"][0]["url"].lower() or "mfa.gov.tr" in external["sources"][0]["url"]
    assert all(
        source["content_hash"] == hashlib.sha256(source["claim"].encode()).hexdigest()
        for source in external["sources"]
    )
    ferry = next(source for source in external["sources"] if source["source_id"] == "sehir-hatlari-timetables")
    assert ferry["verification_state"] == "current-schedule-unverified"

    map_report = json.loads((tmp_path / "map-structure-report.json").read_text())
    assert map_report["editable_overlays"] is True
    fact_checks = json.loads((tmp_path / manifest["fact_check_matrix"]).read_text())
    verdicts = {item["id"]: item["verdict"] for item in fact_checks["checks"]}
    assert verdicts == {
        "visa-mainland-china": "contradicted",
        "ist-m11-old-city": "route-valid-wording-corrected",
        "day2-continent-grouping": "contradicted",
        "ferry-timetable": "unknown-until-travel-date",
        "volatile-prices-and-fees": "author-experience-time-sensitive",
    }
    assert len(map_report["routes"]) == 2
    assert len(map_report["routes"][0]["points"]) == 9

    coverage = json.loads((tmp_path / "content-coverage-matrix.json").read_text())
    assert coverage["claim_count"] == coverage["bound_claim_count"] == 42
    assert coverage["trace_record_count"] == trace["record_count"]
    assert coverage["unbound_claim_count"] == 0
    assert {row["id"] for row in coverage["requirements"]} >= {
        "day1-map",
        "day2-map",
        "lodging",
        "food",
        "price-traps",
        "photo-and-slow-play",
    }

    materials = json.loads((tmp_path / "material-manifest.json").read_text())
    assert "istanbul-osm-d1" in {asset["asset_id"] for asset in materials["assets"]}
    asset_hashes = {
        hashlib.sha256(path.read_bytes()).hexdigest()
        for path in (tmp_path / "assets").iterdir()
        if path.is_file()
    }
    for asset in materials["assets"]:
        assert asset["content_hash"] in asset_hashes

    pptx = Presentation(tmp_path / "istanbul-user-input.pptx")
    pdf = PdfReader(tmp_path / "istanbul-user-input.pdf")
    assert len(pptx.slides) == len(pdf.pages) == 12

    def normalize(value):
        return re.sub(r"[^\w\u4e00-\u9fff]+", "", value or "")

    for index, slide in enumerate(pptx.slides):
        slide_text = "".join(
            shape.text
            for shape in slide.shapes
            if hasattr(shape, "text_frame") and shape.has_text_frame
        )
        assert normalize(slide_text) == normalize(pdf.pages[index].extract_text())
