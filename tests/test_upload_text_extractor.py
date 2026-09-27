import io

from docx import Document

from backend.services.upload_text_extractor import extract_uploaded_text


def test_extracts_text_and_docx_without_persisting_private_upload():
    assert extract_uploaded_text(
        b"private markdown", filename="note.md", content_type="text/markdown",
    ) == "private markdown"
    buffer = io.BytesIO()
    document = Document()
    document.add_heading("Private method", level=1)
    document.add_paragraph("Verified outcome")
    document.save(buffer)
    extracted = extract_uploaded_text(
        buffer.getvalue(), filename="method.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert "Private method" in extracted and "Verified outcome" in extracted


def test_legacy_office_uses_existing_converter_and_removes_temporary_original(monkeypatch):
    from pathlib import Path
    from reportlab.pdfgen import canvas
    import backend.services.presentation_renderer as renderer
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 700, "Hotel reservation address")
    pdf.save()
    paths = []
    def convert(path):
        assert path.read_bytes() == b"legacy fixture"
        paths.append(path)
        return buffer.getvalue()
    monkeypatch.setattr(renderer, "render_office_pdf", convert)
    for suffix in (".doc", ".ppt"):
        assert "Hotel reservation address" in extract_uploaded_text(b"legacy fixture", filename="booking" + suffix, content_type="application/octet-stream")
    assert all(not path.exists() for path in paths)


def test_office_images_join_text_and_deduplicate_without_silent_failure():
    import pytest
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches
    image = io.BytesIO()
    Image.new('RGB', (30, 20), 'blue').save(image, 'PNG')
    doc = Document()
    doc.add_paragraph('Hotel reference')
    doc.add_picture(io.BytesIO(image.getvalue()))
    doc.add_picture(io.BytesIO(image.getvalue()))
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(image.getvalue()), Inches(1), Inches(1))
    for source, suffix in [(doc, '.docx'), (ppt, '.pptx')]:
        buffer = io.BytesIO()
        source.save(buffer)
        seen = []
        def analyze(images):
            seen.extend(images)
            return ['Blue reference image'] * len(images)
        text = extract_uploaded_text(buffer.getvalue(), filename='trip'+suffix, content_type='', analyze_images=analyze)
        assert len(seen) == 1 and 'Blue reference image' in text and 'SHA256' in text
        with pytest.raises(ValueError, match='incomplete'):
            extract_uploaded_text(buffer.getvalue(), filename='trip'+suffix, content_type='', analyze_images=lambda images: [])


def test_image_bridge_cache_is_owner_scoped_and_rejects_untrusted_calls(tmp_path, monkeypatch):
    import asyncio
    import base64
    import httpx
    from PIL import Image
    from scripts import hermes_bridge as bridge
    from scripts.hermes_bridge_runtime import contracts, workflow_artifacts
    monkeypatch.setenv('HERMES_TENANT_SANDBOX_ROOT', str(tmp_path))
    monkeypatch.setattr(contracts, 'HERMES_BRIDGE_INTERNAL_TOKEN', 'vision-test-secret')
    calls = []
    def recognize(goal, node, **kwargs):
        assert kwargs['image_data_urls'][0].startswith('data:image/jpeg;base64,')
        assert kwargs['agent_config'].allowed_tools == []
        calls.append(kwargs['sandbox'].root)
        return 'Test booking 2048', None, {}
    monkeypatch.setattr(workflow_artifacts, '_run_workflow_node_in_process', recognize)
    image = io.BytesIO()
    Image.new('RGB', (20, 20), 'blue').save(image, 'PNG')
    async def run():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=bridge.app), base_url='http://test') as client:
            path = '/v1/documents/analyze-images'
            body = {'images': [base64.b64encode(image.getvalue()).decode()]}
            assert (await client.post(path, json=body)).status_code in (401, 403)
            headers = {'X-Hermes-Internal-Token': 'vision-test-secret', 'X-Tenant-Id': 'tenant', 'X-User-Id': 'alice'}
            for _ in range(2):
                response = await client.post(path, json=body, headers=headers)
                assert response.status_code == 200, response.text
                assert response.json()['analyses'] == ['Test booking 2048']
            assert len(calls) == 1
            response = await client.post(path, json=body, headers={**headers, 'X-User-Id': 'bob'})
            assert response.status_code == 200 and len(calls) == 2 and calls[0] != calls[1]
            assert (await client.post(path, json={'images': ['invalid!']}, headers=headers)).status_code == 422
    asyncio.run(run())


def test_workflow_rejects_provider_error_instead_of_storing_it(monkeypatch, tmp_path):
    import pytest
    import sys
    from types import SimpleNamespace
    from scripts.hermes_bridge_runtime import workflow_artifacts as module, contracts, agent_config, knowledge, memory
    from backend.services.tenant_hermes_sandbox import ensure_tenant_sandbox
    monkeypatch.setenv("HERMES_TENANT_SANDBOX_ROOT", str(tmp_path))
    sandbox = ensure_tenant_sandbox(tenant_key="test-tenant", user_id="test-user")
    fake = SimpleNamespace(run_conversation=lambda _: {"failed": True, "error": "provider unavailable", "final_response": "provider unavailable"}, close=lambda: None)
    monkeypatch.setitem(sys.modules, "run_agent", SimpleNamespace(AIAgent=lambda **_: fake))
    monkeypatch.setitem(sys.modules, "agent.runtime_cwd", SimpleNamespace(set_session_cwd=lambda _: None))
    monkeypatch.setattr(agent_config, "_get_cached_config", lambda: {"model": "test"})
    monkeypatch.setattr(agent_config, "_get_cached_runtime", lambda _: {})
    monkeypatch.setattr(agent_config, "_get_cached_fallback", lambda _: None)
    monkeypatch.setattr(agent_config, "_cache_request_overrides", lambda *_: {})
    monkeypatch.setattr(agent_config, "_create_sandbox_session_db", lambda _: SimpleNamespace(close=lambda: None))
    monkeypatch.setattr(knowledge, "_ensure_tenant_skill_tool_registered", lambda: None)
    monkeypatch.setattr(module, "_workflow_toolsets", lambda *_: [])
    monkeypatch.setattr(memory, "_sandbox_memory_context", lambda _: "")
    with pytest.raises(RuntimeError, match="节点执行失败"):
        module._run_workflow_node_in_process("describe", {"id": "vision", "node_type": "LLM_INFERENCE"}, sandbox=sandbox, agent_config=contracts.TrustedAgentConfig(id="test", allowed_tools=[]))
